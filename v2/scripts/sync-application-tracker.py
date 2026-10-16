"""MCP tracker adapter: fixed labels, exact private bindings, upstream snapshot refresh."""
import json,sys,hashlib,os,tempfile,shutil,copy
from pathlib import Path
from openpyxl import load_workbook

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def run(request):
 config=json.loads(Path(request['configFile']).read_text());book=Path(config['workbook']).resolve();event=request['event'];status=request['applicationStatus']
 if status not in ('待投递','已提交'):raise ValueError('TRACKER_INVALID_LABEL')
 if status!=('已提交' if event['event']=='submission_confirmed' else '待投递'):raise ValueError('TRACKER_EVENT_LABEL_MISMATCH')
 binding=config['roles'][event['companyId']+':'+event['roleId']]
 if binding['roleUrl']!=event['roleUrl'] or binding['title']!=event['roleTitle']:raise ValueError('TRACKER_BINDING_IDENTITY_MISMATCH')
 lock=Path(str(book)+'.application-closeout.lock');fd=os.open(lock,os.O_CREAT|os.O_EXCL|os.O_WRONLY,0o600);tmp=None
 try:
  before=sha(book);wb=load_workbook(book);ws=wb[binding['sheet']];heads=[c.value for c in ws[1]];idx={h:i+1 for i,h in enumerate(heads)}
  if len(heads)!=len(idx) or None in idx:raise ValueError('TRACKER_HEADER_AMBIGUOUS')
  matches=[r for r in range(2,ws.max_row+1) if ws.cell(r,idx['公司']).value==binding['company'] and ws.cell(r,idx['岗位链接']).value==binding['workbookUrl']]
  if len(matches)>1:raise ValueError('TRACKER_DUPLICATE_ROLE')
  created=False
  if not matches:
   fields=binding.get('createFields');
   if not fields or fields.get('公司')!=binding['company'] or fields.get('岗位名称')!=binding['title'] or fields.get('岗位链接')!=binding['workbookUrl']:raise ValueError('TRACKER_ROLE_MISSING')
   if set(fields)-set(heads):raise ValueError('TRACKER_UNKNOWN_COLUMN')
   if any(c.data_type=='f' for sheet in wb for row in sheet for c in row):raise ValueError('TRACKER_INSERT_FORMULAS_REQUIRE_REVIEW')
   same=[r for r in range(2,ws.max_row+1) if ws.cell(r,idx['公司']).value==binding['company']]
   row=max(same)+1 if same else ws.max_row+1;ws.insert_rows(row)
   if same:
    for c in range(1,ws.max_column+1):ws.cell(row,c)._style=copy.copy(ws.cell(same[0],c)._style)
   for k,v in fields.items():ws.cell(row,idx[k],v)
   created=True
  else:row=matches[0]
  if ws.cell(row,idx['岗位名称']).value!=binding['title']:raise ValueError('TRACKER_TITLE_MISMATCH')
  cell=ws.cell(row,idx['投递状态']);old=cell.value
  if old=='已提交' and status!='已提交':raise ValueError('TRACKER_DOWNGRADE_FORBIDDEN')
  if old not in (None,'','待投递','已提交'):raise ValueError('TRACKER_EXISTING_LABEL_INVALID')
  changed=created or old!=status;cell.value=status
  if event['event']!='submission_confirmed':
   marker='[MCP事件 '+event['eventId']+'] '+event['event']+'；证据 '+event['evidence']['file'];note=ws.cell(row,idx['备注']);value=note.value or ''
   if marker not in value:note.value=(value+'\n'+marker).strip();changed=True
  backup=book.parent/'岗位交接'/'备份'/(book.stem+'-closeout-'+before[:12]+'.xlsx');backup.parent.mkdir(parents=True,exist_ok=True)
  if changed:
   if not backup.exists():shutil.copy2(book,backup)
   f,tmp=tempfile.mkstemp(dir=book.parent,suffix='.xlsx');os.close(f)
   for sheet in wb:
    for cells in sheet:
     for cell_item in cells:
      if cell_item.hyperlink:cell_item.hyperlink.ref=cell_item.coordinate
   empty={s.title:{c.coordinate for row in s for c in row if c.value=='' and c.data_type=='s'} for s in wb}
   wb.save(tmp)
   sys.path.insert(0,config['upstreamScripts']);from workbook_store import preserve_empty_text
   preserve_empty_text(tmp,empty)
   if sha(book)!=before:raise ValueError('TRACKER_INPUT_CHANGED')
   os.replace(tmp,book);tmp=None
  wb.close()
  check=load_workbook(book,read_only=True);actual=check[binding['sheet']].cell(row,idx['投递状态']).value;check.close()
  if actual!=status:raise ValueError('TRACKER_READBACK_FAILED')
  # Existing project writer owns full JSON and every registered company-view refresh.
  sys.path.insert(0,config['upstreamScripts']);from workbook_store import refresh,check_full
  refresh_result=refresh(book)
  if check_full(book):raise ValueError('TRACKER_SNAPSHOT_MISMATCH')
  return {'status':'synced','workbook':str(book),'sheet':binding['sheet'],'row':row,'cell':cell.coordinate,'applicationStatus':actual,'createdRow':created,'changed':changed,'workbookHash':sha(book),'backup':str(backup) if changed else None,'fullJson':str(book.parent/'岗位交接'/'工作簿全量.json'),'refresh':refresh_result}
 finally:
  if tmp:Path(tmp).unlink(missing_ok=True)
  os.close(fd);lock.unlink(missing_ok=True)
if __name__=='__main__':
 try:print(json.dumps(run(json.load(sys.stdin)),ensure_ascii=False))
 except Exception as e:print(json.dumps({'status':'sync_failed','error':str(e)},ensure_ascii=False));sys.exit(1)
