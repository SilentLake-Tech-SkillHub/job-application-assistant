import importlib.util,json,tempfile,unittest
from pathlib import Path
from openpyxl import Workbook,load_workbook
spec=importlib.util.spec_from_file_location('tracker_adapter',Path(__file__).parents[1]/'scripts/sync-application-tracker.py');adapter=importlib.util.module_from_spec(spec);spec.loader.exec_module(adapter)
class TrackerSyncTest(unittest.TestCase):
 def setUp(self):
  self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name);self.book=self.root/'book.xlsx';w=Workbook();s=w.active;s.title='Roles';s.append(['公司','岗位名称','岗位链接','投递状态','备注']);s.append(['Sample','Existing','https://example.test/old','待投递','keep']);s.append(['Other','Other','https://example.test/other','待投递','other']);s['C3'].hyperlink='https://example.test/other';w.save(self.book)
  (self.root/'workbook_store.py').write_text("def preserve_empty_text(*args):pass\ndef refresh(book):return {}\ndef check_full(book):return []\n")
  self.binding={'roleUrl':'https://example.test/new','title':'New','sheet':'Roles','company':'Sample','workbookUrl':'https://example.test/new','createFields':{'公司':'Sample','岗位名称':'New','岗位链接':'https://example.test/new'}}
  self.config={'workbook':str(self.book),'upstreamScripts':str(self.root),'roles':{'sample:new':self.binding}}
  self.event={'event':'submission_confirmed','eventId':'test','companyId':'sample','roleId':'new','roleUrl':'https://example.test/new','roleTitle':'New','evidence':{'file':'receipt'}}
 def tearDown(self):
  import sys;sys.modules.pop('workbook_store',None);self.temp.cleanup()
 def call(self):
  f=self.root/'config.json';f.write_text(json.dumps(self.config));return adapter.run({'configFile':str(f),'event':self.event,'applicationStatus':'已提交'})
 def test_insert_preserves_moved_hyperlink_and_repeat_does_not_save(self):
  a=self.call();w=load_workbook(self.book);self.assertEqual(w['Roles']['D3'].value,'已提交');self.assertEqual(w['Roles']['C4'].hyperlink.ref,'C4');self.assertEqual(w['Roles']['C3'].value,'https://example.test/new');self.assertEqual(w['Roles']['E2'].value,'keep');w.close();before=self.book.read_bytes();self.assertFalse(self.call()['changed']);self.assertEqual(self.book.read_bytes(),before)
 def test_missing_role_without_explicit_create_is_rejected(self):
  self.binding.pop('createFields');before=self.book.read_bytes()
  with self.assertRaisesRegex(ValueError,'ROLE_MISSING'):self.call()
  self.assertEqual(before,self.book.read_bytes())
 def test_ambiguous_binding_and_submitted_downgrade_reject(self):
  self.call();self.event['event']='submission_uncertain';f=self.root/'config.json'
  with self.assertRaisesRegex(ValueError,'DOWNGRADE_FORBIDDEN'):adapter.run({'configFile':str(f),'event':self.event,'applicationStatus':'待投递'})
if __name__=='__main__':unittest.main()
