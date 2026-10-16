#!/usr/bin/env python3
"""Stress and robustness for the Excel + 投递证据/ design.

* a large tracker (STRESS_SCALE × 300 companies × 4 roles) with real evidence folders stays fast;
* randomly broken Markdown records and odd cell values never crash the checker or the gates;
* random hook payloads never crash the guard.
Thresholds are generous ceilings for a laptop or CI runner, not performance promises.
"""
import io
import json
import os
import random
import sys
import tempfile
import time
import unittest
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE / 'fixtures'))
sys.path.insert(0, str(HERE.parent / 'hooks'))
from workspace import Workspace  # noqa: E402
from batch_check import Batch  # noqa: E402
from stage_gate import run as gate  # noqa: E402
import application_guard as guard  # noqa: E402

SCALE = max(1, int(os.environ.get('STRESS_SCALE', '1')))
COMPANIES, ROLES = 300 * SCALE, 4


class LargeTrackerTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        names = [('合成压测%04d' % i, '桶%d' % (i % 3)) for i in range(COMPANIES)]
        cls.w = w = Workspace(cls.tmp.name, companies=names, targets='桶0=%d, 桶1=%d, 桶2=无' % (COMPANIES // 3, COMPANIES // 3))
        w.review_enabled = False  # Bulk legacy-state fixtures; handoff exercised for the target below.
        wb = load_workbook(w.book)
        ws, plans = wb['中国大陆'], wb['校招计划']
        statuses = ['已有相关岗位记录', '无用户要求岗位', '无法访问校招页面', '无当年校招计划']
        cls.urls = []
        for i, (company, _) in enumerate(names):
            status = statuses[i % 4] if i % 7 else '已有相关岗位记录'
            for row in plans.iter_rows(min_row=2):
                if row[0].value == company:
                    row[9].value = status
            if status != '已有相关岗位记录':
                continue
            for j in range(ROLES):
                url = 'https://example.test/%d/%d' % (i, j)
                ws.append(['北京', company, '合成', '普通私企', '合成岗位%d' % j, '北京', 'jd', '合成BU', url, '待投递'])
                cls.urls.append((company, url, j))
        # remove the placeholder rows of companies that now have roles
        for row in list(ws.iter_rows(min_row=2)):
            if row[1].value in {c for c, _, _ in cls.urls} and not row[8].value:
                ws.delete_rows(row[0].row)
        wb.save(w.book)
        for company in {c for c, _, _ in cls.urls}:
            w.company_ready(company, limit=ROLES)
        for company, url, j in cls.urls:
            if j == 0:
                w.role_prepared(company, url)
                w.role_approved(company, url, title='合成岗位0', bu='合成BU')
                (w.role_folder(company, url) / '回执.md').write_text('x', encoding='utf-8')
                (w.role_folder(company, url) / '追踪表读回.md').write_text('结果: 通过', encoding='utf-8')
        wb = load_workbook(w.book)
        for row in wb['中国大陆'].iter_rows(min_row=2):
            if row[8].value and str(row[8].value).endswith('/0'):
                row[9].value = '已提交'
        wb.save(w.book)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_consistency_report_and_gates_are_fast(self):
        start = time.perf_counter()
        batch = Batch(self.w.book, 'b1')
        errors = batch.consistency_errors()
        report = batch.report()
        elapsed = time.perf_counter() - start
        self.assertEqual(errors, [])
        self.assertEqual(report['已提交且有回执'], len({c for c, _, _ in self.urls}))
        self.assertLess(elapsed, 60 * SCALE, '%.1fs' % elapsed)
        start = time.perf_counter()
        company, url, _ = next(u for u in self.urls if u[2] == 1)
        self.w.review_enabled = True
        self.w.role_ready(company, url)
        result = gate(self.w.book, 'b1', 'prepare', url)
        self.assertTrue(result['允许'], result['错误'][:3])
        self.assertLess(time.perf_counter() - start, 60 * SCALE)
        print('\n[stress] %d companies / %d role rows: consistency+report %.2fs' % (COMPANIES, len(self.urls), elapsed))


class FuzzTest(unittest.TestCase):
    def test_broken_records_and_cells_never_crash(self):
        rng = random.Random(20261013)
        junk_text = ['', '：', '## ', '依据: ../../x', '依据: C:/x', '- | | |', '目标: 大厂=abc, ,=3', '上限: -1',
                     '结论: ？', '模式: 全自动\n已开启: ,,,', '\x00\x01', '```\n未闭合代码块', '批次: ' + 'x' * 500,
                     '- 简历 | ../outside.pdf | zz', '通道: ' * 30, '审核包哈希: ' + 'a' * 64]
        junk_cells = [None, 0, -5, 3.14, '', ' ', '已提交 ', '尚未推进', '#REF!', 'https://', 'x' * 300, '=1+1']
        with tempfile.TemporaryDirectory() as tmp:
            w = Workspace(tmp)
            w.tag('合成公司A', '已有相关岗位记录'); w.add_role('合成公司A', 'https://example.test/jobs/1')
            w.company_ready('合成公司A'); w.role_prepared('合成公司A', 'https://example.test/jobs/1')
            files = [p for p in Path(tmp, '投递证据').rglob('*.md')]
            originals = {p: p.read_bytes() for p in files}
            for i in range(120):
                for p in rng.sample(files, k=rng.randint(1, 4)):
                    p.write_text(rng.choice(junk_text) + '\n' + rng.choice(junk_text), encoding='utf-8')
                if i % 25 == 0:
                    def mutate(wb):
                        ws = wb[rng.choice(['中国大陆', '校招计划'])]
                        ws.cell(row=rng.randint(2, 6), column=rng.randint(1, 10), value=rng.choice(junk_cells))
                    w.edit(mutate)
                try:
                    batch = Batch(w.book, 'b1')
                    batch.consistency_errors(); batch.closeout(batch.report())
                    for stage in ('query', 'prepare', 'review', 'submit', 'closeout'):
                        gate(w.book, 'b1', stage, 'https://example.test/jobs/1')
                except Exception as exc:  # noqa: BLE001 - any crash is a failure
                    self.fail('iteration %d crashed: %r' % (i, exc))
                for p, data in originals.items():
                    p.write_bytes(data)

    def test_broken_fill_caches_never_crash(self):
        """Random JSON shapes in 填写方案.json: validator, actions, export and every gate stay crash-free."""
        import fill_plan
        rng = random.Random(27)
        atoms = [None, 0, -1, 2.5, '', ' ', '提交申请', '已读回', '投递状态', [], {}, True, 'x' * 300, '档案:', '材料:简历', '用户']
        def junk(depth=0):
            r = rng.random()
            if depth > 2 or r < 0.4:
                return rng.choice(atoms)
            if r < 0.7:
                return [junk(depth + 1) for _ in range(rng.randint(0, 4))]
            keys = list(fill_plan.TOP_KEYS | fill_plan.FIELD_KEYS | {'投递状态', '回执', 'label', 'css', 'role', '点击', '选项', '等待'})
            return {rng.choice(keys): junk(depth + 1) for _ in range(rng.randint(0, 6))}
        with tempfile.TemporaryDirectory() as tmp:
            w = Workspace(tmp)
            url = 'https://example.test/jobs/1'
            w.tag('合成公司A', '已有相关岗位记录'); w.add_role('合成公司A', url)
            w.company_ready('合成公司A'); w.role_prepared('合成公司A', url)
            path = w.role_folder('合成公司A', url) / fill_plan.PLAN_NAME
            good = json.loads(path.read_text(encoding='utf-8'))
            for i in range(3000):
                data = json.loads(json.dumps(good))
                if rng.random() < 0.5:
                    data = junk()
                else:
                    target = rng.choice([data] + [f for f in data['字段']])
                    target[rng.choice(list(target) + ['投递状态', '定位'])] = junk()
                try:
                    errors = fill_plan.validate(data)
                    fill_plan.validate(data, ready=True); fill_plan.validate_template(data)
                    if not errors:
                        fill_plan.actions(data); fill_plan.pending(data)
                    if i % 30 == 0:
                        path.write_text(json.dumps(data, ensure_ascii=False) if rng.random() < 0.9 else '{broken', encoding='utf-8')
                        batch = Batch(w.book, 'b1')
                        batch.consistency_errors(); batch.report()
                        for stage in ('prepare', 'review', 'submit'):
                            gate(w.book, 'b1', stage, url)
                except Exception as exc:  # noqa: BLE001 - any crash is a failure
                    self.fail('iteration %d crashed: %r on %s' % (i, exc, json.dumps(data, ensure_ascii=False)[:300]))

    def test_hook_payloads_never_crash(self):
        rng = random.Random(7)
        junk = [None, 0, 1.5, '', [], {}, '提交', 'Submit', {'action': 'left_click'}, ['x', {'a': '提交'}],
                {'actions': 'x'}, {'actions': [None, 1, {'name': None, 'input': '提交'}]}]
        with tempfile.TemporaryDirectory() as tmp:
            Path(tmp, '投递证据').mkdir()
            old = sys.stderr
            try:
                for _ in range(20000):
                    payload = {'hook_event_name': rng.choice(['PreToolUse', 'Stop', None, 1]), 'cwd': rng.choice([tmp, '/nonexistent', None, 5]),
                               'tool_name': rng.choice(['mcp__x__computer', 'mcp__playwright__browser_click', 'Bash', '', None, 3]),
                               'tool_input': rng.choice(junk)}
                    if rng.random() < 0.3:
                        payload = rng.choice(junk)
                    sys.stderr = io.StringIO()
                    self.assertIn(guard.main(json.dumps(payload, ensure_ascii=False)), (0, 2))
            finally:
                sys.stderr = old


if __name__ == '__main__':
    unittest.main()
