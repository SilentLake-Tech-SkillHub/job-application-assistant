"""Synthetic company review: no live company, account, application or user approval."""
import copy
import json
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path
from openpyxl import Workbook, load_workbook
import company_handoff as h


def sample_company(book,name='合成公司',key='synthetic'):
 wb=load_workbook(book);roles=[]
 for ws in wb:
  if ws.title in {v[0] for v in h.SCHEMA.values()}:continue
  heads=[x.value for x in ws[1]]
  if '岗位链接' not in heads:continue
  for values in ws.iter_rows(min_row=2,values_only=True):
   row=dict(zip(heads,values))
   if row.get('公司')!=name or not row.get('岗位链接'):continue
   roles.append({'id':'r'+str(len(roles)+1),'plan_ids':['regular'],'title':row['岗位名称'],'url':row['岗位链接'],'bu':row['所属BU'],'cities':row['工作地点'],'employment':'正式','tier':'exact','functional_fit':'符合','jd':row['岗位JD'],'basis':'合成需求及JD匹配','eligibility':'符合','eligibility_basis':'合成材料','recommended':True,'selected':True,'order':len(roles)+1,'source':'https://example.test/jobs','checked_at':'20XX-10-16'})
 wb.close()
 return {'company_key':key,'name':name,'version':'v1','scope':'合成测试：AI产品、城市不限','coverage':'complete','gaps':'','checked_at':'20XX-10-16','recommendation_reason':'基于合成JD和名额建议','plans':[{'id':'regular','name':'普通校招','type':'校招','requirements':'毕业窗口与学历，见合成来源','window':'本招聘轮次','source':'https://example.test/campus','checked_at':'20XX-10-16','special':False,'decision':'consider','decision_ref':''}], 'roles':roles,'rules':[{'id':dim,'dimension':dim,'scope':'全部岗位','status':'verified','text':'合成规则；无额外限制；公司上限单列','limit':len(roles)+5 if dim=='company' else '', 'used':0 if dim=='company' else '', 'role_ids':[], 'shared_group':'','source':'https://example.test/faq','checked_at':'20XX-10-16','gap':''} for dim in sorted(h.DIMENSIONS)],'referrals':[{'id':'ref','scope':'全部计划/BU','role_ids':[],'status':'not_found','code':'','source':'https://example.test/referral','checked_at':'20XX-10-16','expires':'','notes':'已查公开来源，无可核实码'}]}


class CompanyHandoffTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.book=self.root/'tracker.xlsx';self.batch='synthetic'
  wb=Workbook();ws=wb.active;ws.title='地区岗位';ws.append(['公司','岗位名称','所属BU','工作地点','岗位链接','岗位JD','投递状态']);ws.append(['合成公司','AI产品','BU-A','城市甲','https://example.test/job/1','职责：AI产品需求。要求：本科。','待投递']);ws.append(['合成公司','用户产品','BU-B','城市乙','https://example.test/job/2','职责：用户产品。要求：本科。','待投递']);wb.save(self.book)
  self.c=sample_company(self.book);self.c['roles'][1].update(tier='partial',selected=False,order='',recommended=True)
 def tearDown(self):self.tmp.cleanup()
 def draft(self):h.write(self.book,self.batch,{'companies':[self.c]})
 def approve(self):
  self.draft();(self.root/'confirmation.md').write_text('合成测试确认当前HTML的选择与全部规则；不是真实用户授权', encoding='utf-8')
  h.confirm(self.book,self.batch,self.c['company_key'],'confirmation.md',h.sha(h.paths(self.book,self.batch)/'公司核对.html'))
 def test_legacy_default_encoding_json_first_roundtrip(self):
  original_read,original_write=Path.read_text,Path.write_text
  def legacy_read(path,encoding=None,**kwargs):return original_read(path,encoding=encoding or 'cp1252',**kwargs)
  def legacy_write(path,data,encoding=None,**kwargs):return original_write(path,data,encoding=encoding or 'cp1252',**kwargs)
  with patch.object(Path,'read_text',legacy_read),patch.object(Path,'write_text',legacy_write):
   draft=self.root/'draft.json';draft.write_text(json.dumps({'companies':[self.c]},ensure_ascii=False),encoding='utf-8')
   with patch('sys.argv',['company_handoff.py','write','--workbook',str(self.book),'--batch',self.batch,'--input',str(draft)]):
    self.assertEqual(h.main(),0)
   (self.root/'confirmation.md').write_text('合成用户确认',encoding='utf-8')
   h.confirm(self.book,self.batch,self.c['company_key'],'confirmation.md',h.sha(h.paths(self.book,self.batch)/'公司核对.html'))
   self.assertEqual(h.check(self.book,self.batch,self.c['roles'][0]['url']),[])
 def test_full_flow_and_separate_selection(self):
  self.approve();self.assertEqual(h.check(self.book,self.batch,self.c['roles'][0]['url']),[])
  self.assertIn('岗位不在用户最终选择中',h.check(self.book,self.batch,self.c['roles'][1]['url']))
 def test_unconfirmed(self):
  self.draft();self.assertTrue(h.check(self.book,self.batch,self.c['roles'][0]['url']))
 def test_json_tamper(self):
  self.approve();p=h.paths(self.book,self.batch)/'公司方案.json';d=json.loads(p.read_text(encoding='utf-8'));d['companies'][0]['roles'][1]['selected']=True;p.write_text(json.dumps(d), encoding='utf-8');self.assertIn('JSON内容与Excel不一致',h.check(self.book,self.batch))
 def test_excel_refresh_stale_then_export(self):
  self.approve();wb=load_workbook(self.book);wb['地区岗位']['G2']='已提交';wb.save(self.book)
  self.assertIn('Excel/JSON未同步',h.check(self.book,self.batch));h.export(self.book,self.batch);self.assertEqual(h.check(self.book,self.batch),[])
 def test_business_change_invalidates_confirmation(self):
  self.approve();wb=load_workbook(self.book);ws=wb['投递规则'];ws['G2']='changed official rule';wb.save(self.book);h.export(self.book,self.batch)
  self.assertIn('公司方案未确认或确认已失效',h.check(self.book,self.batch))
 def test_html_tamper(self):
  self.approve();p=h.paths(self.book,self.batch)/'公司核对.html';p.write_text('changed', encoding='utf-8');self.assertIn('HTML与Excel不一致',h.check(self.book,self.batch))
 def test_confirmation_tamper(self):
  self.approve();(self.root/'confirmation.md').write_text('changed', encoding='utf-8');self.assertIn('用户确认文件变化',h.check(self.book,self.batch))
 def test_wrong_html_hash(self):
  self.draft();(self.root/'confirmation.md').write_text('synthetic', encoding='utf-8');
  with self.assertRaises(ValueError):h.confirm(self.book,self.batch,'synthetic','confirmation.md','wrong')
 def test_missing_rule_dimension(self):
  self.c['rules']=self.c['rules'][:-1]
  with self.assertRaises(ValueError):self.draft()
 def test_per_bu_scope_cannot_omit_other_candidates(self):
  next(x for x in self.c['rules'] if x['dimension']=='bu')['role_ids']=['r1']
  with self.assertRaises(ValueError):self.draft()
 def test_missing_referral(self):
  self.c['referrals']=[]
  with self.assertRaises(ValueError):self.draft()
 def test_per_scope_referral(self):
  self.c['referrals'][0]['role_ids']=['r1']
  with self.assertRaises(ValueError):self.draft()
 def test_special_pending(self):
  self.c['plans'][0].update(special=True,decision='pending')
  with self.assertRaises(ValueError):self.draft()
 def test_special_accepted(self):
  (self.root/'plan-confirm.md').write_text('synthetic: consider special plan', encoding='utf-8');self.c['plans'][0].update(special=True,decision='consider',decision_ref='plan-confirm.md');self.c['roles'][0]['tier']='special';self.approve();self.assertEqual(h.check(self.book,self.batch),[])
 def test_quota_each_dimension(self):
  for dim in ('company','round','bu','shared','plan','concurrent'):
   with self.subTest(dim=dim):
    c=copy.deepcopy(self.c);next(x for x in c['rules'] if x['dimension']==dim).update(limit=0,role_ids=['r1'])
    with self.assertRaises(ValueError):h.write(self.book,self.batch,{'companies':[c]})
 def test_missing_regional_row(self):
  self.approve();wb=load_workbook(self.book);wb['地区岗位'].delete_rows(2);wb.save(self.book);h.export(self.book,self.batch);self.assertIn('地区岗位表缺少唯一对应行',h.check(self.book,self.batch))
 def test_regional_jd_changed(self):
  self.approve();wb=load_workbook(self.book);wb['地区岗位']['F2']='新要求';wb.save(self.book);h.export(self.book,self.batch);self.assertIn('地区表与确认目录不一致: 岗位JD',h.check(self.book,self.batch))
 def test_outside_directory_and_escaping(self):
  r=copy.deepcopy(self.c['roles'][1]);r.update(id='r3',title='<script>alert(1)</script>',tier='outside',jd='',selected=False,recommended=False,order='',url='',source='javascript:alert(1)');self.c['roles'].append(r);self.draft()
  page=(h.paths(self.book,self.batch)/'公司核对.html').read_text(encoding='utf-8');self.assertIn('&lt;script&gt;',page);self.assertNotIn('href="javascript:',page);self.assertIn('目录岗位 3 个',page)
 def test_preserve_other_sheets_and_formula_like_jd(self):
  self.c['roles'][1]['jd']='=HYPERLINK("x")';self.draft();wb=load_workbook(self.book);self.assertEqual(wb['地区岗位']['G2'].value,'待投递');self.assertEqual(wb['岗位目录']['L3'].data_type,'s')
 def test_long_text_rejected_not_truncated(self):
  self.c['roles'][0]['jd']='中'*32768
  with self.assertRaises(ValueError):self.draft()
 def test_same_version_refresh_rejected(self):
  self.draft()
  with self.assertRaises(ValueError):self.draft()
 def test_new_version_requires_new_confirmation(self):
  self.approve();self.c['version']='v2';self.draft();self.assertIn('公司方案未确认或确认已失效',h.check(self.book,self.batch))

 def test_empty_confirmation_rejected(self):
  self.draft();(self.root/'empty.md').write_text('', encoding='utf-8')
  with self.assertRaises(ValueError):h.confirm(self.book,self.batch,'synthetic','empty.md',h.sha(h.paths(self.book,self.batch)/'公司核对.html'))
 def test_duplicate_role_url_rejected(self):
  self.c['roles'][1]['url']=self.c['roles'][0]['url']
  with self.assertRaises(ValueError):self.draft()
 def test_expired_referral_blocks_handoff(self):
  self.c['referrals'][0].update(status='verified',code='TEST',expires='2000-01-01');self.approve()
  self.assertIn('内推码已过期，须复核并同步',h.check(self.book,self.batch))
 def test_missing_json_after_interrupted_export(self):
  self.approve();(h.paths(self.book,self.batch)/'公司方案.json').unlink();self.assertTrue(h.check(self.book,self.batch));h.export(self.book,self.batch);self.assertEqual(h.check(self.book,self.batch),[])
 def test_refresh_keeps_previous_versions(self):
  self.approve();old=h.sha(h.paths(self.book,self.batch)/'公司方案.json');self.c['version']='v2';self.draft()
  self.assertTrue((h.paths(self.book,self.batch)/'历史'/(old+'.json')).is_file())

if __name__=='__main__':unittest.main()
