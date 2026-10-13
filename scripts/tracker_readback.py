#!/usr/bin/env python3
"""Read back the saved role tracker after a write (workbook contract).

Reopens the saved .xlsx and compares it with the pre-write backup:
  * the backup is the file read before writing (its SHA-256 matches --pre-sha256);
  * every expected row exists exactly once by its key and shows the expected values;
  * other rows in the touched sheets, untouched sheets, formulas and sheet order are unchanged;
  * no cell shows an Excel error value.
Structural readback only; it does not prove a website accepted an application.

  tracker_readback.py --workbook T.xlsx --backup T.bak.xlsx --pre-sha256 <hex> \
      --expect "中国大陆|岗位链接|https://…/jobs/101|投递状态=已提交" \
      --expect "校招计划|公司|合成公司|盘点状态=已有相关岗位记录" \
      [--record 投递证据/<批次>/岗位/<岗位目录>/追踪表读回.md]

Each --expect is  工作表|键列|键值|列=值[;列=值…].
"""
import argparse
import hashlib
import json
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - reported as a readable failure
    load_workbook = None

ERROR_VALUES = {'#REF!', '#VALUE!', '#NAME?', '#DIV/0!', '#N/A', '#NUM!', '#NULL!'}


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def cell(value):
    return '' if value is None else str(value).strip()


def table(ws, key_column, keyless=None):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return None, {}, ['sheet ' + ws.title + ' is empty']
    header = [cell(h) for h in rows[0]]
    if key_column not in header:
        return header, {}, ['key column ' + repr(key_column) + ' not in header of ' + ws.title]
    k = header.index(key_column)
    by_key = {}
    for number, row in enumerate(rows[1:], start=2):
        values = dict(zip(header, (cell(v) for v in row)))
        key = cell(row[k]) if k < len(row) else ''
        if key:
            by_key.setdefault(key, []).append((number, values))
        elif keyless is not None and any(values.values()):
            keyless.append(values)
    return header, by_key, []


def sheet_values(ws):
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


def parse_expect(text):
    parts = text.split('|')
    if len(parts) < 4:
        raise ValueError('--expect needs 工作表|键列|键值|列=值: ' + text)
    sheet, key_column, key = parts[0].strip(), parts[1].strip(), parts[2].strip()
    fields = {}
    for pair in '|'.join(parts[3:]).split(';'):
        if pair.strip():
            column, _, wanted = pair.partition('=')
            fields[column.strip()] = wanted.strip()
    return {'sheet': sheet, 'key_column': key_column, 'key': key, 'fields': fields}


def check(workbook, backup, pre_sha256, expectations):
    if load_workbook is None:
        return ['openpyxl is required: pip install openpyxl']
    if not Path(workbook).is_file():
        return ['workbook not found']
    if not Path(backup).is_file():
        return ['backup not found: make a recoverable backup before writing']
    errors = []
    if sha256(backup) != pre_sha256.lower():
        errors.append('backup SHA-256 differs from the hash recorded before writing')
    if sha256(workbook) == sha256(backup) and expectations:
        errors.append('workbook is byte-identical to the backup: the write was not saved')
    after_f, before_f = load_workbook(workbook), load_workbook(backup)
    after_v = load_workbook(workbook, data_only=True)
    if after_f.sheetnames != before_f.sheetnames:
        errors.append('sheet list or order changed')
    touched = {}
    for item in expectations:
        touched.setdefault((item['sheet'], item['key_column']), []).append(item)
    touched_sheets = {sheet for sheet, _ in touched}
    for name in after_f.sheetnames:
        if name not in touched_sheets and name in before_f.sheetnames and sheet_values(after_f[name]) != sheet_values(before_f[name]):
            errors.append('other sheet changed: ' + name)
        for row in after_v[name].iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) and v.strip() in ERROR_VALUES:
                    errors.append('error value ' + v + ' in sheet ' + name)
    for (sheet, key_column), items in touched.items():
        if sheet not in after_f.sheetnames:
            errors.append('sheet ' + repr(sheet) + ' not found')
            continue
        after_keyless, before_keyless = [], []
        header, after_rows, problems = table(after_f[sheet], key_column, after_keyless)
        errors.extend(problems)
        if header is None or problems:
            continue
        _, before_rows, _ = table(before_f[sheet], key_column, before_keyless) if sheet in before_f.sheetnames else (None, {}, [])
        # Rows without a key (company placeholders) must stay as they were; a placeholder may only
        # disappear when it became that company's first expected role row.
        expected_companies = {after_rows[i['key']][0][1].get('公司') for i in items if len(after_rows.get(i['key'], [])) == 1}
        remaining = list(after_keyless)
        for row in before_keyless:
            if row in remaining:
                remaining.remove(row)
            elif row.get('公司') not in expected_companies:
                errors.append('%s: placeholder row changed or removed: %s' % (sheet, row.get('公司') or row))
        for row in remaining:
            errors.append('%s: unexpected new or changed row without key: %s' % (sheet, row.get('公司') or row))
        keys = set()
        for item in items:
            keys.add(item['key'])
            matches = after_rows.get(item['key'], [])
            if len(matches) != 1:
                errors.append('%s: expected exactly one row for %s, found %d' % (sheet, item['key'], len(matches)))
                continue
            number, row = matches[0]
            for column, wanted in item['fields'].items():
                if column not in header:
                    errors.append('%s: column %r not in header' % (sheet, column))
                elif row.get(column) != wanted:
                    errors.append('%s row %d %s reads back %r, expected %r' % (sheet, number, column, row.get(column), wanted))
        for key, rows in before_rows.items():
            if key not in keys and [r for _, r in after_rows.get(key, [])] != [r for _, r in rows]:
                errors.append('%s: unrelated row changed or removed: %s' % (sheet, key))
        for key in set(after_rows) - set(before_rows) - keys:
            errors.append('%s: unexpected new row: %s' % (sheet, key))
        for row_f, row_b in zip(after_f[sheet].iter_rows(), before_f[sheet].iter_rows()):
            for a, b in zip(row_f, row_b):
                if isinstance(b.value, str) and b.value.startswith('=') and a.value != b.value:
                    errors.append('formula changed at %s!%s' % (sheet, a.coordinate))
    return errors


def record(path, workbook, pre_sha256, expectations, errors):
    lines = ['# 追踪表读回', '', '结果: ' + ('通过' if not errors else '未通过'),
             '写入前sha256: ' + pre_sha256, '写入后sha256: ' + sha256(workbook), '工作簿: ' + Path(workbook).name, '', '## 核对行']
    for item in expectations:
        lines.append('- ' + ' | '.join([item['sheet'], item['key_column'], item['key'],
                                         '; '.join(k + '=' + v for k, v in item['fields'].items())]))
    if errors:
        lines += ['', '## 问题'] + ['- ' + e for e in errors]
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text('\n'.join(lines) + '\n', encoding='utf-8')


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--workbook', required=True, type=Path)
    parser.add_argument('--backup', required=True, type=Path)
    parser.add_argument('--pre-sha256', required=True)
    parser.add_argument('--expect', action='append', default=[], required=True)
    parser.add_argument('--record', type=Path, help='write the readback result as Markdown evidence')
    args = parser.parse_args()
    try:
        expectations = [parse_expect(e) for e in args.expect]
        errors = check(args.workbook, args.backup, args.pre_sha256, expectations)
    except (OSError, ValueError, KeyError) as exc:
        print(json.dumps({'读回通过': False, '错误': [str(exc)]}, ensure_ascii=False))
        return 2
    if args.record:
        record(args.record, args.workbook, args.pre_sha256, expectations, errors)
    print(json.dumps({'读回通过': not errors, '错误': errors, '核对行数': len(expectations),
                      '工作簿sha256': sha256(args.workbook)}, ensure_ascii=False, indent=2))
    return 0 if not errors else 1


if __name__ == '__main__':
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding='utf-8')
        except (AttributeError, ValueError):
            pass
    sys.exit(main())
