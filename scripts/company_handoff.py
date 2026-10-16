#!/usr/bin/env python3
"""Excel-owned company review; derived JSON/HTML and evidence-bound application handoff."""
import argparse
import copy
import hashlib
import html
import json
import os
import shutil
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit
from structured_dimensions import EXTRA_FIELDS, LIST_FIELDS as DIMENSION_LISTS, NULL_FIELDS, applies, validate as validate_dimensions
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SCHEMA = {
 'company': ('公司核对', [('company_key','公司键'),('name','公司'),('version','版本'),('scope','本批需求'),('coverage','覆盖状态'),('gaps','覆盖缺口'),('checked_at','核验时间'),('recommendation_reason','推荐理由'),('confirmation_ref','确认依据'),('confirmation_sha256','确认文件哈希'),('confirmed_at','确认时间'),('confirmed_digest','确认数据摘要'),('reviewed_html_sha256','已审核HTML哈希')]),
 'plans': ('计划明细', [('id','计划ID'),('name','计划名称'),('type','计划类型'),('requirements','计划要求'),('window','开放窗口'),('source','来源'),('checked_at','核验时间'),('special','特殊计划'),('decision','用户决定'),('decision_ref','决定依据')]),
 'roles': ('岗位目录', [('id','岗位ID'),('plan_ids','所属计划ID'),('title','岗位名称'),('url','岗位链接'),('bu','所属BU'),('cities','工作地点'),('employment','用工性质'),('tier','候选档位'),('functional_fit','职责符合程度'),('jd','完整JD'),('basis','分档依据'),('eligibility','资格结论'),('eligibility_basis','资格依据'),('recommended','Agent推荐'),('selected','用户选择'),('order','用户次序'),('source','来源'),('checked_at','核验时间')]),
 'rules': ('投递规则', [('id','规则ID'),('dimension','规则维度'),('scope','适用范围'),('status','核验状态'),('text','规则原文及解释'),('limit','上限'),('used','已用名额'),('role_ids','适用岗位ID'),('shared_group','共享组'),('source','来源'),('checked_at','核验时间'),('gap','未知项')]),
 'referrals': ('内推信息', [('id','记录ID'),('scope','适用范围'),('role_ids','适用岗位ID'),('status','核验状态'),('code','内推码'),('source','来源'),('checked_at','核验时间'),('expires','有效期'),('notes','说明')]),
}
SCHEMA.update({
 'conditions': ('条件明细', [('id','条件ID'),('subject_type','关联类型'),('subject_id','关联ID'),('field','条件维度'),('operator','比较关系'),('condition_value','条件值'),('unit','单位'),('group_id','条件组'),('group_operator','组内关系'),('importance','条件性质'),('status','核验状态'),('source','来源'),('checked_at','核验时间'),('gap','未知项'),('raw','原文依据')]),
 'preferences': ('志愿明细', [('id','志愿ID'),('role_id','岗位ID'),('group_id','志愿组'),('rank','志愿序号'),('label','志愿名称'),('mode','志愿机制'),('bu','志愿BU'),('cities','志愿地点'),('status','状态'),('source','来源'),('checked_at','核验时间')]),
 'research': ('研究检查', [('id','检查项'),('status','完成状态'),('source','来源'),('checked_at','核验时间'),('gap','待补事项'),('expected_count','应有数量'),('observed_count','实际数量')]),
})
for kind, fields in EXTRA_FIELDS.items():SCHEMA[kind][1].extend(fields)
RESEARCH_ITEMS = {'plans','catalog','jd','eligibility','rules','referrals','preferences'}
DIMENSIONS = {'company','plan','round','bu','shared','concurrent','cooldown','change','preference'}
CONFIRM_FIELDS = {'confirmation_ref','confirmation_sha256','confirmed_at','confirmed_digest','reviewed_html_sha256'}
LIST_FIELDS = {'plan_ids','role_ids'}|DIMENSION_LISTS
JSON_FIELDS={'condition_value'}
BOOL_FIELDS = {'special','recommended','selected'}

def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def digest(obj): return hashlib.sha256(json.dumps(obj, ensure_ascii=False, sort_keys=True, separators=(',',':')).encode()).hexdigest()
def now(): return datetime.now(timezone.utc).isoformat(timespec='seconds')
def paths(book,batch):
 if not batch or Path(batch).name != batch or batch in ('.','..'): raise ValueError('无效批次')
 return Path(book).resolve().parent/'岗位交接'/batch

def atomic(path, data):
 path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
 fd,name=tempfile.mkstemp(dir=path.parent,suffix='.tmp');os.close(fd)
 try:
  Path(name).write_bytes(data);os.replace(name,path)
 finally:
  Path(name).unlink(missing_ok=True)

def encode(k,v):
 if k in JSON_FIELDS:return json.dumps(v,ensure_ascii=False)
 if k in NULL_FIELDS and v is None:return None
 if k in LIST_FIELDS:return json.dumps(v or [],ensure_ascii=False)
 if k in BOOL_FIELDS:return bool(v)
 return v if v not in (None,'') else None
def decode(k,v):
 if k in JSON_FIELDS:return json.loads(v) if v not in (None,'') else None
 if k in NULL_FIELDS:return None if v in (None,'') else v
 if k in LIST_FIELDS:return json.loads(v or '[]')
 if k in BOOL_FIELDS:
  if v not in (True,False,None):raise ValueError('布尔列取值无效: '+k)
  return v is True
 return v if v is not None else ''
def headers(kind):
 return ['批次']+(['公司键'] if kind!='company' else [])+[v for _,v in SCHEMA[kind][1]]

def read(book,batch,allow_legacy=False):
 wb=load_workbook(book,data_only=False);companies={}
 for kind,(sheet,fields) in SCHEMA.items():
  if sheet not in wb:
   if allow_legacy and kind in ('preferences','research','conditions'):continue
   raise ValueError('缺少公司交接工作表: '+sheet)
  ws=wb[sheet]
  actual=[c.value for c in ws[1]]
  if actual!=headers(kind) and not (allow_legacy and actual==headers(kind)[:len(actual)]):raise ValueError('交接表头不一致: '+sheet)
  for row in ws.iter_rows(min_row=2,values_only=True):
   if row[0]!=batch:continue
   offset=1 if kind=='company' else 2
   data={k:decode(k,v) for (k,_),v in zip(fields,row[offset:])}
   key=data['company_key'] if kind=='company' else row[1]
   if kind=='company':
    if key in companies:raise ValueError('公司键重复')
    companies[key]={**data,**{k:[] for k in SCHEMA if k!='company'}}
   else:
    if key not in companies:raise ValueError('孤立的明细公司键')
    companies[key][kind].append(data)
 # Index original rows once; avoid rescanning the whole workbook per company.
 by_name={}
 for c in companies.values():
  c['existing_data']=[];by_name.setdefault(c['name'],[]).append(c)
 for ws in wb:
  if ws.title in {x[0] for x in SCHEMA.values()}:continue
  names=[x.value for x in ws[1]]
  if '公司' not in names:continue
  for vals in ws.iter_rows(min_row=2,values_only=True):
   row=dict(zip(names,vals))
   for c in by_name.get(row.get('公司'),[]):
    c['existing_data'].append({'sheet':ws.title,'values':{str(k):v.isoformat() if hasattr(v,'isoformat') else v for k,v in row.items() if k is not None}})
 wb.close();return list(companies.values())

def business(c):
 return {k:v for k,v in c.items() if k not in CONFIRM_FIELDS|{'existing_data'}}

def save(book,batch,companies,updates=None):
 book=Path(book);before=sha(book);wb=load_workbook(book)
 from workbook_store import apply_updates
 apply_updates(wb,updates or [],book.resolve().parent)
 keys={c['company_key'] for c in companies}
 if len(keys)!=len(companies):raise ValueError('公司键重复')
 for kind,(sheet,fields) in SCHEMA.items():
  ws=wb[sheet] if sheet in wb else wb.create_sheet(sheet)
  if ws.max_row==1 and ws['A1'].value is None:ws.append(headers(kind));ws.delete_rows(1)
  actual=[c.value for c in ws[1]]
  if actual!=headers(kind):
   if actual!=headers(kind)[:len(actual)]:raise ValueError('已有同名工作表结构不同: '+sheet)
   for column,title in enumerate(headers(kind)[len(actual):],len(actual)+1):ws.cell(1,column,title)
  for i in range(ws.max_row,1,-1):
   if ws.cell(i,1).value==batch and ws.cell(i,2).value in keys:ws.delete_rows(i)
  for c in companies:
   for item in ([c] if kind=='company' else c[kind]):
    vals=[batch]+([c['company_key']] if kind!='company' else [])+[encode(k,item.get(k,'')) for k,_ in fields]
    if any(isinstance(v,str) and len(v)>32767 for v in vals):raise ValueError('Excel单元格过长；禁止截断')
    ws.append(vals)
    for cell,v in zip(ws[ws.max_row],vals):
     if isinstance(v,str):cell.data_type='s'  # never interpret employer text as formulas
  ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
  for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='173B56')
  for row in ws.iter_rows(min_row=2):
   for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
  for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=28
 backup=book.parent/'岗位交接'/'备份'/(book.stem+'-'+before[:12]+'.xlsx');backup.parent.mkdir(parents=True,exist_ok=True)
 if not backup.exists():shutil.copy2(book,backup)
 fd,tmp=tempfile.mkstemp(dir=book.parent,suffix='.xlsx');os.close(fd)
 try:
  empty_text={ws.title:{c.coordinate for row in ws for c in row if c.value=='' and c.data_type=='s'} for ws in wb}
  wb.save(tmp);wb.close()
  from workbook_store import preserve_empty_text
  preserve_empty_text(tmp,empty_text)
  if sha(book)!=before:raise ValueError('工作簿已被其他进程修改')
  os.replace(tmp,book)
 finally:Path(tmp).unlink(missing_ok=True)

def safe_ref(root,ref):
 p=(root/str(ref)).resolve()
 if not ref or not p.is_relative_to(root.resolve()) or not p.is_file() or p.stat().st_size==0:raise ValueError('缺少有效私有证据文件: '+str(ref))
 return p

def validate_company(c):
 errors=[]
 for kind,(_,fields) in SCHEMA.items():
  for item in ([c] if kind=='company' else c.get(kind,[])):
   for key,_ in fields:
    if key in BOOL_FIELDS and type(item.get(key))!=bool:errors.append('布尔维度须为true/false: '+key)
 def need(obj,keys,label):
  for k in keys:
   if not obj.get(k):errors.append(label+' 缺少 '+k)
 need(c,['company_key','name','version','scope','coverage','recommendation_reason'],'公司')
 if c.get('coverage') not in ('complete','partial'):errors.append('覆盖状态必须 complete/partial')
 if c.get('coverage')=='partial' and not c.get('gaps'):errors.append('覆盖不完整必须列缺口')
 if c.get('coverage')=='complete' and c.get('gaps'):errors.append('有缺口不可称完整覆盖')
 if c.get('coverage')=='complete' and not c.get('checked_at'):errors.append('完整覆盖缺核验时间')
 plans={p['id']:p for p in c['plans']};ids={r['id'] for r in c['roles']}
 urls=[r['url'] for r in c['roles'] if r['url']]
 if len(urls)!=len(set(urls)):errors.append('同一岗位直链重复，须合并计划关系')
 for section in ('plans','roles','rules','referrals','preferences','research','conditions'):
  values=[x.get('id') for x in c[section]]
  if len(values)!=len(set(values)):errors.append(section+' ID重复')
 for p in c['plans']:
  need(p,['id','name','type','requirements','window','source','checked_at'],'计划')
  if p['special'] and p['decision'] not in ('consider','decline','pending'):errors.append('特殊计划缺用户决定')
  if p['special'] and p['decision'] in ('consider','decline') and not p['decision_ref']:errors.append('特殊计划决定缺依据')
 for r in c['roles']:
  need(r,['id','title','plan_ids','bu','cities','employment','tier','source','checked_at'],'目录岗位')
  if not set(r['plan_ids'])<=plans.keys():errors.append('岗位计划不存在')
  if r['tier'] not in ('exact','partial','special','outside','pending'):errors.append('候选档位无效')
  if r['tier'] in ('exact','partial','special') or r['selected'] or r['recommended']:
   need(r,['url','jd','basis','eligibility','eligibility_basis'],'候选')
  if r['selected'] or r['recommended']:
   if r['tier'] in ('outside','pending'):errors.append('不可选择/推荐范围外或待核候选')
   if not r['url'].startswith(('https://','http://')):errors.append('岗位URL无效')
   if r['selected'] and (type(r['order'])!=int or r['order']<1):errors.append('已选岗位缺顺序')
   for pid in r['plan_ids']:
    p=plans.get(pid,{})
    if p.get('special') and p.get('decision')!='consider':errors.append('特殊计划未经考虑确认')
 orders=[r['order'] for r in c['roles'] if r['selected']]
 if len(orders)!=len(set(orders)):errors.append('选岗顺序重复')
 dimensions={r['dimension'] for r in c['rules']}
 if not DIMENSIONS<=dimensions:errors.append('规则维度缺失: '+','.join(sorted(DIMENSIONS-dimensions)))
 for role in c['roles']:
  if role['tier'] in ('exact','partial','special'):
   for dimension in DIMENSIONS:
    if not any(rule['dimension']==dimension and applies(rule,role) for rule in c['rules']):errors.append('候选缺适用规则: '+role['id']+'/'+dimension)
 for rule in c['rules']:
  need(rule,['id','scope','status','text'],'规则')
  if rule['status']!='unknown':need(rule,['source','checked_at'],'规则')
  if rule['status'] not in ('verified','unknown','not_applicable'):errors.append('规则状态无效')
  if rule['status']=='unknown' and not rule['gap']:errors.append('未知规则缺具体缺口')
  if not set(rule['role_ids'])<=ids:errors.append('规则引用未知岗位')
  if rule['limit']!='':
   if type(rule['limit'])!=int or rule['limit']<0:errors.append('规则上限必须非负整数');continue
   if rule['status']!='verified':errors.append('未核实规则不可填确定上限')
   if rule['used']!='' and (type(rule['used'])!=int or rule['used']<0):errors.append('已用名额无效')
   used=rule['used'] if type(rule['used'])==int else 0
   for flag in ('selected','recommended'):
    count=sum(bool(r[flag]) and applies(rule,r) for r in c['roles'])
    if rule['dimension']=='preference':
     chosen={r['id'] for r in c['roles'] if r[flag] and applies(rule,r)}
     count=sum(v['role_id'] in chosen and v['status']=='selected' and (not rule['shared_group'] or v['group_id']==rule['shared_group']) for v in c['preferences'])
    if count+used>rule['limit']:errors.append(flag+' 超过规则 '+rule['id'])
 if not c['referrals']:errors.append('缺少内推核查')
 for ref in c['referrals']:
  need(ref,['id','scope','status'],'内推')
  if ref['status']!='unverified':need(ref,['source','checked_at'],'内推')
  elif not ref.get('notes'):errors.append('未核实内推缺待补说明')
  if ref['status'] not in ('verified','not_found','unverified','not_applicable'):errors.append('内推状态无效')
  if ref['status']=='verified' and not ref['code']:errors.append('已核内推缺码')
  if not set(ref['role_ids'])<=ids:errors.append('内推引用未知岗位')
 for r in c['roles']:
  if r['tier'] in ('exact','partial','special') and not any(applies(x,r) for x in c['referrals']):errors.append('候选缺对应内推核查: '+r['id'])
 errors+=validate_research(c)+validate_dimensions(c)
 return errors


def validate_research(c,ready=False):
 errors=validate_dimensions(c,True) if ready else [];roles={r['id']:r for r in c['roles']}
 tasks={r['id']:r for r in c.get('research',[])}
 if set(tasks)!=RESEARCH_ITEMS:errors.append('研究检查必须覆盖计划、目录、JD、资格、规则、内推、志愿')
 for task in tasks.values():
  status=task.get('status')
  if status not in ('todo','checked','blocked'):errors.append('研究检查状态无效')
  if status=='checked' and (not task.get('source') or not task.get('checked_at')):errors.append('已查项目缺来源或时间')
  if status!='checked' and not task.get('gap'):errors.append('未完成检查缺待补事项')
  if (ready or c.get('coverage')=='complete') and status!='checked':errors.append('研究未完成: '+task['id'])
  for key in ('expected_count','observed_count'):
   if task.get(key,'')!='' and (type(task[key])!=int or task[key]<0):errors.append('研究数量必须非负整数')
  if task['id'] in ('plans','catalog') and status=='checked':
   expected=task.get('expected_count');observed=task.get('observed_count')
   actual=len(c['plans'] if task['id']=='plans' else c['roles'])
   if type(expected)!=int or observed!=actual or expected!=observed:errors.append('计划/目录数量未对齐: '+task['id'])
 pref=c.get('preferences',[]);slots=set();per_role={}
 for v in pref:
  rid=v.get('role_id');per_role.setdefault(rid,[]).append(v)
  if rid not in roles:errors.append('志愿引用未知岗位')
  if v.get('status') not in ('selected','pending_selection','not_applicable','unknown'):errors.append('志愿状态无效')
  if v.get('mode') not in ('ordered','parallel','not_applicable','unknown'):errors.append('志愿机制无效')
  if v.get('status')=='selected':
   if not v.get('group_id') or not v.get('label') or v.get('mode') not in ('ordered','parallel'):errors.append('已定志愿缺组、名称或机制')
   if v.get('mode')=='ordered':
    if type(v.get('rank'))!=int or v['rank']<1:errors.append('顺序志愿缺正整数序号')
    slot=(v.get('group_id'),v.get('rank'))
    if slot in slots:errors.append('同一志愿组序号重复')
    slots.add(slot)
   elif v.get('rank') is not None:errors.append('平行志愿不填顺序')
  elif v.get('rank') is not None:errors.append('未选/无志愿/未知志愿的序号必须为空')
  if v.get('status')=='not_applicable' and v.get('mode')!='not_applicable':errors.append('无志愿机制与状态不一致')
  if v.get('mode')!='unknown' and rid in roles:
   for rule in c['rules']:
    if rule['dimension']=='preference' and rule['status']=='verified' and applies(rule,roles[rid]) and rule.get('preference_mode')!=v.get('mode'):errors.append('志愿记录与规则机制不一致')
  if v.get('status')!='unknown' and (not v.get('source') or not v.get('checked_at')):errors.append('志愿缺来源或时间')
  if ready and rid in roles and roles[rid]['selected'] and v.get('status') in ('unknown','pending_selection'):errors.append('已选岗位志愿待核实')
 for rid,role in roles.items():
  if role['tier'] in ('exact','partial','special') and rid not in per_role:errors.append('候选缺志愿记录: '+rid)
  values=per_role.get(rid,[])
  if len(values)>1 and any(v['status']!='selected' for v in values):errors.append('同岗位志愿记录互相冲突')
 return errors

def preference_target(c,url):
 ids={r['id'] for r in c['roles'] if r['url']==url}
 values=[v for v in c['preferences'] if v['role_id'] in ids]
 if values and all(v['status']=='not_applicable' for v in values):return '无'
 return '；'.join(v['group_id']+'/'+('parallel' if v['mode']=='parallel' else str(v['rank']))+'/'+v['label']+'/'+v.get('bu','')+'/'+v.get('cities','') for v in sorted(values,key=lambda v:(v['group_id'],v['rank'] or 0)))

def audit(book,batch):
 return [{'company_key':c['company_key'],'company':c['name'],'issues':list(dict.fromkeys(validate_company(c)+validate_research(c,True)))} for c in read(book,batch)]

def render(companies,batch):
 esc=lambda x:html.escape(str(x if x is not None else ''))
 def val(x):
  if x is None:return ''
  if isinstance(x,list):return '、'.join(map(str,x))
  if x is True:return '是'
  if x is False:return '否'
  labels={'exact':'完全符合','partial':'部分符合','special':'特殊计划','outside':'范围外目录记录','pending':'待确认','complete':'覆盖完整','consider':'考虑','decline':'不考虑','verified':'已核实','unknown':'未核实','not_applicable':'不适用','not_found':'未找到','unverified':'未核实','company':'公司总量','plan':'计划限额','round':'投递轮次','bu':'BU限额','shared':'共享名额','concurrent':'同时流程','cooldown':'冷冻期','change':'撤回与换岗','preference':'志愿规则','todo':'待查','checked':'已查','blocked':'受阻','selected':'已定志愿','ordered':'有序志愿','parallel':'平行志愿','pending_selection':'待用户选择','regular':'普通校招','technical':'技术计划','product_trainee':'产品培训生','management_trainee':'管理培训生','known':'已知','not_disclosed':'未披露','structured':'已结构化','count':'数量限制','unlimited':'不限','independent':'独立','exclusive':'互斥','conditional':'有条件','none':'无'}
  return labels.get(str(x),str(x))
 def table(rows,fields):
  out='<table><thead><tr>'+''.join('<th>'+esc(title)+'</th>' for _,title in fields)+'</tr></thead><tbody>'
  for r in rows:
   out+='<tr>'
   for key,_ in fields:
    x=val(r.get(key,''));cell=esc(x)
    if key in ('url','source') and urlsplit(x).scheme in ('http','https'):cell='<a href="'+esc(x)+'">'+cell+'</a>'
    out+='<td>'+cell+'</td>'
   out+='</tr>'
  return out+'</tbody></table>'
 def records(rows,fields,title_key):
  return ''.join('<details open><summary>'+esc(val(row.get(title_key,'')))+'</summary>'+table([{'label':title,'value':val(row.get(key,''))} for key,title in fields],[('label','项目'),('value','内容')])+'</details>' for row in rows)
 parts=['<!doctype html><html lang="zh"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>公司投递方案核对</title><style>body{font:16px/1.7 system-ui;background:#f2f5f8;color:#172c3b;margin:0}main{max-width:1200px;margin:auto;padding:32px}section{background:white;padding:28px;margin:24px 0;border-radius:16px}h1,h2,h3{color:#173b56}table{border-collapse:collapse;width:100%;margin:18px 0;font-size:14px}td,th{border:1px solid #dae2e8;padding:10px;text-align:left;vertical-align:top;white-space:pre-wrap;overflow-wrap:anywhere}th{background:#e9f0f5}pre{white-space:pre-wrap;overflow-wrap:anywhere}a{color:#16638a}.note{background:#fff3ce;padding:16px}details{margin:14px 0}summary{cursor:pointer;font-weight:600}</style><main><h1>公司投递方案 · 用户核对版</h1><p>批次 '+esc(batch)+' · Excel 为主记录；此页用于核对，确认不等于最终提交批准。</p>']
 for c in companies:
  parts+=['<section><h2>'+esc(c['name'])+' · '+esc(c['version'])+'</h2><p>'+esc(c['scope'])+'</p><p>已发现计划 '+str(len(c['plans']))+' 个；目录岗位 '+str(len(c['roles']))+' 个；覆盖 '+esc('覆盖完整' if c['coverage']=='complete' else '覆盖未完整')+'；核验 '+esc(c['checked_at'])+'</p><p class="note">缺口：'+esc(c['gaps'] or '本次记录无缺口；仍需核对原始来源')+'</p>']
  parts+=['<h3>需求与公司范围</h3>',table([{'label':title,'value':val(c.get(key,''))} for key,title in SCHEMA['company'][1] if key not in CONFIRM_FIELDS],[('label','项目'),('value','内容')]),'<h3>可选计划与要求</h3>',records(c['plans'],SCHEMA['plans'][1],'name'),'<h3>计划与岗位条件明细</h3>',table(c['conditions'],SCHEMA['conditions'][1]),'<h3>全部岗位名称目录</h3>',table(c['roles'],[(k,t) for k,t in SCHEMA['roles'][1] if k in ('id','plan_ids','round_id','title','url','bu','cities','city_options','employment','tier')])]
  for tier,title in [('exact','完全符合'),('partial','部分符合'),('special','特殊计划')]:
   parts+=['<h3>'+title+'</h3>']
   for r in c['roles']:
    if r['tier']==tier:parts+=['<details open><summary>'+esc(r['title'])+'</summary>'+table([{'label':t,'value':val(r[k])} for k,t in SCHEMA['roles'][1] if k!='jd'],[('label','项目'),('value','内容')])+'<h4>完整 JD</h4><pre>'+esc(r['jd'])+'</pre></details>']
  parts+=['<h3>最终推荐</h3><p>'+esc(c['recommendation_reason'])+'</p>',table([r for r in c['roles'] if r['recommended']],[('title','岗位'),('bu','BU'),('cities','地点'),('basis','理由')]),'<h3>用户最终选择（待当前版本明确确认）</h3>',table(sorted([r for r in c['roles'] if r['selected']],key=lambda r:r['order']),[('order','次序'),('title','岗位'),('bu','BU'),('cities','地点'),('url','链接')]),'<h3>投递轮次、公司及 BU 数量与冷冻期</h3>',records(c['rules'],SCHEMA['rules'][1],'dimension'),'<h3>内推码与适用范围</h3>',table(c['referrals'],SCHEMA['referrals'][1]),'<h3>志愿（与处理次序独立）</h3>',table(c['preferences'],SCHEMA['preferences'][1]),'<h3>研究完成与待补检查</h3>',table(c['research'],SCHEMA['research'][1]),'<h3>现有 Excel 公司记录</h3>']
  for old in c.get('existing_data',[]):parts+=['<details><summary>'+esc(old['sheet'])+'</summary><pre>'+esc(json.dumps(old['values'],ensure_ascii=False,indent=2))+'</pre></details>']
  parts+=['</section>']
 return ''.join(parts)+'</main></html>'

def export_batch(book,batch,full):
 out=paths(book,batch);source_sha=sha(book);companies=read(book,batch)
 if source_sha!=full['workbook_sha256']:raise ValueError('完整导出期间工作簿变化')
 errors=[e for c in companies for e in validate_company(c)]
 if errors:raise ValueError('; '.join(errors))
 payload={'schema_version':2,'batch_id':batch,'workbook_sha256':source_sha,'full_data_sha256':full['data_sha256'],'companies':companies}
 payload['data_sha256']=digest(companies)
 if sha(book)!=source_sha:raise ValueError('导出期间工作簿变化，须重读')
 for name in ('公司方案.json','公司核对.html'):
  old=out/name
  if old.exists():atomic(out/'历史'/(sha(old)+old.suffix),old.read_bytes())
 atomic(out/'公司方案.json',json.dumps(payload,ensure_ascii=False,indent=2).encode())
 atomic(out/'公司核对.html',render(companies,batch).encode())
 return payload

def export(book,batch):
 from workbook_store import refresh
 refresh(book,batch)
 return json.loads((paths(book,batch)/'公司方案.json').read_text(encoding='utf-8'))

def write(book,batch,data):
 json.dumps(data,ensure_ascii=False,allow_nan=False)
 if set(data)-{'companies','company_patches','workbook_updates'}:raise ValueError('输入含未知顶层字段')
 if 'company_patches' in data:
  if 'companies' in data:raise ValueError('companies与company_patches只能选一种')
  current={c['company_key']:c for c in read(book,batch)};items=[]
  for patch in data['company_patches']:
   key=patch['company_key']
   if key not in current:raise ValueError('增量更新公司不存在，先init')
   if not patch.get('version'):raise ValueError('增量更新须提供新版本')
   c=copy.deepcopy(current[key])
   for kind,changes in patch.items():
    if kind in SCHEMA and kind!='company':
     rows={x['id']:x for x in c[kind]}
     for change in changes:
      rid=change['id'];rows.setdefault(rid,{}).update(change)
     c[kind]=list(rows.values())
    else:c[kind]=changes
   items.append(c)
 else:items=copy.deepcopy(data.get('companies',[]))
 for c in items:
  allowed={k for k,_ in SCHEMA['company'][1]}|set(SCHEMA)-{'company'}|{'existing_data'}
  if set(c)-allowed:raise ValueError('公司含未知字段')
  for kind,(_,fields) in SCHEMA.items():
   if kind!='company':
    for item in c.get(kind,[]):
     if set(item)-{k for k,_ in fields}:raise ValueError(kind+' 含未知字段')
  for k in CONFIRM_FIELDS:c[k]=''
  for kind,(_,fields) in SCHEMA.items():
   for item in ([c] if kind=='company' else c.setdefault(kind,[])):
    for k,_ in fields:item.setdefault(k,None if k in NULL_FIELDS|JSON_FIELDS else [] if k in LIST_FIELDS else False if k in BOOL_FIELDS else '')
  errors=validate_company(c)
  if errors:raise ValueError('; '.join(errors))
 try:old={c['company_key']:c for c in read(book,batch,allow_legacy=True)}
 except ValueError as exc:
  if '缺少公司交接工作表' not in str(exc):raise
  old={}
 for c in items:
  if c['company_key'] in old and c['version']==old[c['company_key']]['version']:raise ValueError('刷新须使用新版本')
 save(book,batch,items,data.get('workbook_updates'));return export(book,batch)

def confirm(book,batch,key,evidence,html_sha):
 out=paths(book,batch);root=Path(book).resolve().parent
 snapshot=json.loads((out/'公司方案.json').read_text(encoding='utf-8'))
 if snapshot['workbook_sha256']!=sha(book):raise ValueError('Excel刷新后须重新导出核对')
 if sha(out/'公司核对.html')!=html_sha:raise ValueError('用户审核HTML版本不一致')
 companies=read(book,batch)
 if digest(companies)!=snapshot['data_sha256'] or snapshot['companies']!=companies:raise ValueError('JSON与Excel不一致')
 if (out/'公司核对.html').read_text(encoding='utf-8')!=render(companies,batch):raise ValueError('HTML被修改')
 matches=[c for c in companies if c['company_key']==key]
 if len(matches)!=1:raise ValueError('公司身份不唯一')
 c=matches[0]
 errors=validate_company(c)+validate_research(c,True)
 if errors:raise ValueError('; '.join(errors))
 from workbook_store import check_full
 if check_full(book):raise ValueError('完整JSON未同步')
 ref=safe_ref(root,evidence)
 archive=out/'历史'/(html_sha+'.html');atomic(archive,(out/'公司核对.html').read_bytes())
 c.update(confirmation_ref=evidence,confirmation_sha256=sha(ref),confirmed_at=now(),confirmed_digest=digest(business(c)),reviewed_html_sha256=html_sha)
 save(book,batch,[c]);return export(book,batch)

def check(book,batch,url=None,target=None):
 try:
  out=paths(book,batch);root=Path(book).resolve().parent
  payload=json.loads((out/'公司方案.json').read_text(encoding='utf-8'));companies=read(book,batch)
  from workbook_store import check_full,full_path
  errors=check_full(book)
  full=json.loads(full_path(book).read_text(encoding='utf-8'))
  if payload.get('full_data_sha256')!=full.get('data_sha256'):errors.append('批次JSON与全量JSON版本不一致')
  if payload.get('schema_version')!=2 or payload.get('batch_id')!=batch:errors.append('交接协议或批次不一致')
  if payload.get('workbook_sha256')!=sha(book):errors.append('Excel/JSON未同步')
  if payload.get('companies')!=companies or payload.get('data_sha256')!=digest(companies):errors.append('JSON内容与Excel不一致')
  if not (out/'公司核对.html').is_file() or (out/'公司核对.html').read_text(encoding='utf-8')!=render(companies,batch):errors.append('HTML与Excel不一致')
  targets=[c for c in companies if url is None or any(r['url']==url for r in c['roles'])]
  if not targets or (url and len(targets)!=1):errors.append('交接中没有唯一对应岗位')
  for c in targets:
   errors+=validate_company(c)+validate_research(c,True)
   if target is not None and target!=preference_target(c,url):errors.append('目标确认的志愿与公司方案不一致')
   for ref in c['referrals']:
    if ref['status']=='verified' and ref['expires']:
     try:
      if datetime.fromisoformat(ref['expires']).date()<datetime.now(timezone.utc).date():errors.append('内推码已过期，须复核并同步')
     except ValueError:errors.append('内推有效期需ISO日期或留空并说明未知')
   if not c['confirmed_at'] or c['confirmed_digest']!=digest(business(c)):errors.append('公司方案未确认或确认已失效')
   try:
    if sha(safe_ref(root,c['confirmation_ref']))!=c['confirmation_sha256']:errors.append('用户确认文件变化')
    if sha(out/'历史'/(c['reviewed_html_sha256']+'.html'))!=c['reviewed_html_sha256']:errors.append('已确认HTML变化')
   except (ValueError,OSError):errors.append('缺少用户确认或已确认HTML证据')
   for p in c['plans']:
    if p['special'] and p['decision'] in ('consider','decline'):
     try:safe_ref(root,p['decision_ref'])
     except ValueError:errors.append('特殊计划决定证据不存在')
   for r in c['roles']:
    if url and r['url']!=url:continue
    if url and not r['selected']:errors.append('岗位不在用户最终选择中')
    if not r['selected']:continue
    rows=[x['values'] for x in c['existing_data'] if x['values'].get('岗位链接')==r['url']]
    if len(rows)!=1:errors.append('地区岗位表缺少唯一对应行');continue
    for key,col in [('title','岗位名称'),('bu','所属BU'),('cities','工作地点'),('jd','岗位JD')]:
     if rows[0].get(col)!=r[key]:errors.append('地区表与确认目录不一致: '+col)
  return list(dict.fromkeys(errors))
 except (OSError,ValueError,KeyError,TypeError) as exc:return ['交接不可读取: '+str(exc)]

def main():
 for stream in (sys.stdout,sys.stderr):
  if hasattr(stream,'reconfigure'):stream.reconfigure(encoding='utf-8',errors='backslashreplace')
 p=argparse.ArgumentParser(description=__doc__);p.add_argument('command',choices=['write','export','confirm','check','audit']);p.add_argument('--workbook',required=True,type=Path);p.add_argument('--batch',required=True);p.add_argument('--input',type=Path);p.add_argument('--company');p.add_argument('--evidence');p.add_argument('--html-sha256');p.add_argument('--url');a=p.parse_args()
 try:
  if a.command=='write':write(a.workbook,a.batch,json.loads(a.input.read_text(encoding='utf-8')))
  elif a.command=='export':export(a.workbook,a.batch)
  elif a.command=='audit':
   print(json.dumps(audit(a.workbook,a.batch),ensure_ascii=False,indent=2));return 0
  elif a.command=='confirm':confirm(a.workbook,a.batch,a.company,a.evidence,a.html_sha256)
  else:
   errors=check(a.workbook,a.batch,a.url)
   print(json.dumps({'通过':not errors,'错误':errors},ensure_ascii=False));return int(bool(errors))
  print('已保存、读回并同步');return 0
 except Exception as exc:print(str(exc));return 1
if __name__=='__main__':raise SystemExit(main())
