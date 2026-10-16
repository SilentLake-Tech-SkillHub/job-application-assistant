#!/usr/bin/env python3
"""Full-flow test: one complete batch, Query → Prepare → Review → Submit → Resume → Closeout.

State lives only in the Excel tracker (the campus sample workbook) and 投递证据/; there is no JSON
ledger. The run uses the synthetic recruiting site in fixtures/synthetic_ats.py with a real headless
browser (in-app lane), the real stage gates, the PreToolUse submit guard and tracker readback.

Batch: 合成公司 (two roles on the synthetic site, limit 2) and 合成公司B (campus page unreachable),
补位规则 留空 (keep_gap). The Chrome extension is "missing", so 合成公司 runs in the Playwright backup lane:
its forms are filled by scripts/playwright_fill.py from a 填写方案.json built from the control
inventory, with one run interrupted and resumed from the cache. User decisions are simulated and stored as 确认记录 files. A pass proves the
Skill's procedure and tooling fit together end to end, not that any live ATS behaves the same way.

  python3 test_e2e_full_flow.py [output_dir]
"""
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE / 'fixtures'))
from synthetic_ats import serve  # noqa: E402
from workspace import Workspace, today, now  # noqa: E402
from stage_gate import run as gate  # noqa: E402
from batch_check import Batch  # noqa: E402
from workbook_state import Workbook  # noqa: E402
import fill_plan  # noqa: E402
from test_playwright_fill import plan_for  # noqa: E402
from browser_session import profile_endpoint  # noqa: E402

A, B = '合成公司', '合成公司B'
ENV = dict(os.environ, PYTHONIOENCODING='utf-8')


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class FullFlow(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.out = Path(sys.argv[1]) if len(sys.argv) > 1 and not sys.argv[1].startswith('-') else None
        cls.tmp = tempfile.TemporaryDirectory()
        cls.w = Workspace(cls.tmp.name, batch='e2e', companies=((A, '全部'), (B, '全部')),
                          targets='全部=无', rule='留空', basis='有合格岗位的公司')
        cls.w.write_query('全部=无', '留空', '有合格岗位的公司',
                          能力预检='内置浏览器: 可用\nChrome: 不可用\ncomputer use: 未检查\n系统脚本: 不适用\nPlaywright: 可用')
        cls.profile = cls.w.root / '.浏览器配置' / 'playwright'
        cls.base, cls.site, cls.server = serve()
        cls.pw = sync_playwright().start()
        cls.browser = cls.pw.chromium.launch(headless=True, executable_path=os.environ.get('CHROME_EXECUTABLE') or None)
        cls.page = cls.browser.new_page(viewport={'width': 1000, 'height': 900})
        cls.log = []

    @classmethod
    def tearDownClass(cls):
        endpoint = profile_endpoint(cls.profile)
        if endpoint:
            # Explicit cleanup of this test's own temporary retained browser.
            retained = cls.pw.chromium.connect_over_cdp(endpoint)
            retained.new_browser_cdp_session().send('Browser.close')
            for _ in range(50):
                if profile_endpoint(cls.profile) is None:
                    break
                time.sleep(0.1)
        cls.browser.close(); cls.pw.stop(); cls.server.shutdown(); cls.server.server_close()
        if cls.out:
            cls.out.mkdir(parents=True, exist_ok=True)
            shutil.copytree(cls.w.root / '投递证据', cls.out / '投递证据', dirs_exist_ok=True)
            shutil.copy(cls.w.book, cls.out / cls.w.book.name)
            shutil.copytree(cls.w.root / '岗位交接', cls.out / '岗位交接', dirs_exist_ok=True)
            (cls.out / 'flow-log.json').write_text(json.dumps(cls.log, ensure_ascii=False, indent=2), encoding='utf-8')
        cls.tmp.cleanup()

    # ------------------------------------------------------------ helpers
    def gate(self, stage, url=None, allowed=True):
        result = gate(self.w.book, 'e2e', stage, url)
        self.log.append({'gate': stage, 'url': url, 'allowed': result['允许'], 'errors': result['错误'][:4]})
        self.assertEqual(result['允许'], allowed, json.dumps(result, ensure_ascii=False, indent=1))
        return result

    def cli_gate(self, stage, url):
        result = subprocess.run([sys.executable, str(HERE / 'stage_gate.py'), '--workbook', str(self.w.book), '--batch', 'e2e',
                                 '--stage', stage, '--url', url], capture_output=True, text=True, encoding='utf-8', env=ENV)
        self.log.append({'gate': stage + ' (CLI)', 'url': url, 'allowed': result.returncode == 0})
        return result

    def hook(self, summary):
        payload = json.dumps({'hook_event_name': 'PreToolUse', 'session_id': 'e2e', 'cwd': str(self.w.root),
                              'tool_name': 'mcp__Claude_Browser__computer',
                              'tool_input': {'action': 'left_click', 'action_summary': summary}}, ensure_ascii=False)
        result = subprocess.run([sys.executable, str(HERE.parent / 'hooks' / 'application_guard.py')],
                                input=payload.encode('utf-8'), capture_output=True, env=ENV)
        self.log.append({'hook': summary, 'exit': result.returncode})
        return result.returncode

    def shot(self, folder, name):
        path = Path(folder) / name
        path.parent.mkdir(parents=True, exist_ok=True)
        self.page.screenshot(path=str(path), full_page=True)
        return path

    def note(self, folder, name, lines):
        path = Path(folder) / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text('\n'.join(lines) + '\n', encoding='utf-8')
        return path

    def write_tracker(self, record, expectations, mutate):
        """Reread, back up, smallest edit, save, reopen and read back (workbook contract)."""
        pre = sha(self.w.book)
        backup = self.w.root / ('追踪表备份-%d.xlsx' % len(self.log))
        shutil.copy(self.w.book, backup)
        self.w.edit(mutate)
        args = [sys.executable, str(HERE / 'tracker_readback.py'), '--workbook', str(self.w.book), '--backup', str(backup),
                '--pre-sha256', pre, '--record', str(record)]
        for e in expectations:
            args += ['--expect', e]
        result = subprocess.run(args, capture_output=True, text=True, encoding='utf-8', env=ENV)
        self.assertEqual(result.returncode, 0, result.stdout)

    # ------------------------------------------------------------ the flow
    def test_full_flow(self):
        self.m1_query()
        self.m2_scan_and_tag()
        self.m2_prepare()
        self.m3_review()
        self.m4_submit()
        self.m5_resume()
        self.m6_closeout()

    def m1_query(self):
        self.gate('query')  # the fixture wrote Query.md with every must-ask section and its user record
        query = self.w.bdir / '批次' / 'Query.md'
        saved = query.read_text(encoding='utf-8')
        query.write_text(saved.replace('## 自动化授权', '## 自动化授权（未问）'), encoding='utf-8')
        self.gate('query', allowed=False)  # a missing must-ask item blocks the batch
        query.write_text(saved, encoding='utf-8')

    def m2_scan_and_tag(self):
        page, w = self.page, self.w
        # 合成公司B: the campus page is unreachable → scanned, tagged, explained in 备注 (keep_gap).
        page.goto(self.base + '/no-such-campus-page')
        b_dir = w.root / '投递证据' / 'e2e' / '公司' / B
        self.shot(b_dir, '访问失败.png')
        self.gate('closeout', allowed=False)
        self.assertTrue(any('合成公司B 已经扫描过' in e for e in Batch(w.book, 'e2e').consistency_errors()))
        w.tag(B, '无法访问校招页面', '2026-10-13 e2e 批次：校招页面 404，留空不补位；证据 投递证据/e2e/公司/合成公司B/访问失败.png')
        # 合成公司: log in, read the count rule and every in-scope role, then tag and write rows.
        page.goto(self.base + '/center')
        page.get_by_role('button', name='登录').click()  # browser-prefilled account; password never read
        page.wait_for_url(self.base + '/center')
        self.assertIn('synthetic-user', page.locator('#account').inner_text())
        a_dir = w.root / '投递证据' / 'e2e' / '公司' / A
        rule = page.locator('#rule').inner_text()
        self.assertIn('最多投递 2 个岗位', rule)
        self.shot(a_dir, '规则.png')
        self.note(a_dir, '数量规则.md', ['状态: 已核实', '上限: 2', '来源: ' + self.base + '/center', '核验日期: ' + today(), '原文: ' + rule])
        self.note(a_dir, '通道.md', ['通道: Playwright', '说明: Chrome 插件不可用，Playwright 作备用通道'])
        self.urls = {}
        for job in ('101', '102'):
            page.goto(self.base + '/jobs/' + job)
            title, bu, city, jd = (page.locator(s).inner_text() for s in ('#title', '#bu', '#city', '#jd'))
            url = page.url
            self.urls[job] = dict(url=url, title=title, bu=bu, city=city, jd=jd)
            w.add_role(A, url, title=title, bu=bu, city=city, jd=jd)
            folder = w.role_folder(A, url)
            self.note(folder, '岗位包.md', ['公司: ' + A, 'BU: ' + bu, '办公地: ' + city, '岗位: ' + title, '直链: ' + url, '完整JD: ' + jd])
            self.assertIn('必需：本科及以上', jd)
            self.note(folder, '资格核对.md', ['结论: 符合', '必需: 本科及以上 → 简历学历：硕士'])
        w.tag(A, '已有相关岗位记录')
        self.assertEqual(Batch(w.book, 'e2e').consistency_errors(), [])
        # Company order not confirmed yet: forms stay closed.
        self.gate('prepare', self.urls['101']['url'], allowed=False)
        self.note(a_dir, '次序确认.md', ['决定: 不变', '依据: ' + w.say('order', '合成公司两个岗位次序不变')])
        w.confirm_company_review(A, limit=2)
        self.page.goto((w.root / '岗位交接' / 'e2e' / '公司核对.html').as_uri())
        self.assertIn('合成公司', self.page.locator('body').inner_text())
        self.shot(a_dir, '公司HTML核对.png')

    def m2_prepare(self):
        self.login_profile()
        for job, slot in (('101', '合成智能事业群'), ('102', '合成数据平台部')):
            url = self.urls[job]['url']
            self.gate('prepare', url)
            self.fill(job, slot)
            self.gate('review', url, allowed=False)  # no review package yet

    def login_profile(self):
        """The user logs in once in the Playwright profile; the executor reuses that login."""
        from browser_session import connect
        browser, ctx, _ = connect(self.pw, self.profile, headless=True,
                                  executable_path=os.environ.get('CHROME_EXECUTABLE') or None)
        self.profile_browser = browser
        page = ctx.new_page(); page.goto(self.base + '/login')
        page.get_by_role('button', name='登录').click(); page.wait_for_url(self.base + '/center')
        # Keep the dedicated synthetic session alive, as the current executor does.


    def executor(self, plan_path, *extra):
        before = sha(self.w.book)
        result = subprocess.run([sys.executable, str(HERE / 'playwright_fill.py'), '--plan', str(plan_path), '--headless',
                                 '--profile-dir', str(self.profile), '--workbook', str(self.w.book), '--timeout', '4000', *extra],
                                capture_output=True, text=True, encoding='utf-8', env=ENV)
        self.assertEqual(sha(self.w.book), before, 'Filling must never write Excel state')
        self.log.append({'executor': Path(plan_path).parent.name, 'exit': result.returncode})
        return result.returncode, json.loads(result.stdout)

    def fill(self, job, slot):
        page, w = self.page, self.w
        url = self.urls[job]['url']
        folder = w.role_folder(A, url)
        page.goto(self.base + '/apply/' + job); page.wait_for_selector('#save')
        page.add_script_tag(path=str(HERE / 'inspect_form_controls.js'))
        inventory = page.evaluate('inspectApplicationForm(document)')
        self.assertTrue({'file', 'select', 'checkbox'} <= {c['kind'] for c in inventory['fields']})
        self.note(folder, '控件盘点.md', ['控件数: %d' % inventory['count']] + ['- %s | %s' % (c['label'], c['kind']) for c in inventory['fields']])
        # Skeleton from the inventory, then the Agent fills in values and the real click paths.
        plan = fill_plan.from_inventory(inventory, 'e2e', A, url, 'Playwright')
        plan['表单链接'] = self.base + '/apply/' + job
        known = {f['标签']: f for f in plan_for(self.base, job)['字段']}
        known['第一意向部门 *']['值'] = slot  # user-choice slot, confirmed later in 目标确认
        for f in plan['字段']:
            if f['标签'] in known:
                keep, label = f['序号'], f['标签']
                f.clear(); f.update(known[label]); f['序号'] = keep
            elif not f.get('必填'):
                f['状态'] = '跳过'; f['备注'] = '选填，本次不填'
        plan['保存步骤'] = [{'点击': {'role': 'button', 'name': '保存草稿'}}]
        path = folder / fill_plan.PLAN_NAME
        fill_plan.save(path, plan)
        self.assertEqual(fill_plan.validate(plan, ready=True), [])
        self.gate('review', url, allowed=False)  # cache exists but nothing read back yet
        if job == '101':
            code, out = self.executor(path, '--max-fields', '2')  # interrupted
            self.assertEqual(code, 0, out)
            cached = fill_plan.load(path)['字段']
            done = sum(1 for f in cached if f['状态'] in fill_plan.DONE)
            self.assertEqual(sum(1 for f in cached if f['状态'] == '已读回'), 2)
            self.assertIn('%d/%d 已读回' % (done, len(cached)), Batch(w.book, 'e2e').report()['填写进度'].values())
        code, out = self.executor(path)
        self.assertEqual(code, 0, out)
        self.assertIn('已点击保存', out['保存'])
        fill_plan.export_readback(path)
        page.goto(self.base + '/apply/' + job); page.wait_for_selector('#save')
        reopened = self.readback()
        self.assertEqual(reopened, {'姓名': '合成候选人', '学历': '硕士', '毕业时间': '2026-06',
                                    '项目经历': '合成项目：负责 AI 产品需求、方案与迭代，指标提升 20%。', '第一意向': slot,
                                    '条款': 'True', '附件': '简历.pdf'})
        self.shot(folder, '草稿.png')
        self.note(folder, '草稿读回.md', ['%s: %s' % kv for kv in reopened.items()])

    def readback(self):
        p = self.page
        return {'姓名': p.get_by_label('姓名 *').input_value(), '学历': p.locator('.combo input').input_value(),
                '毕业时间': p.get_by_label('毕业时间 *').input_value(), '项目经历': p.get_by_label('项目经历 *').input_value(),
                '第一意向': p.get_by_label('第一意向部门 *').input_value(),
                '条款': str(p.get_by_label('我已阅读并同意《合成隐私条款》 *').is_checked()),
                '附件': p.locator('#resume-name').inner_text()}

    def m3_review(self):
        for job in ('101', '102'):
            url = self.urls[job]['url']
            folder = self.w.role_folder(A, url)
            answers = (folder / '草稿读回.md').read_text(encoding='utf-8')
            self.note(folder, '审核包.md', ['# 审核包', '材料: 简历.pdf', '账号名额: 0/2 已用', '未决字段: 无', '截图: 填写.png', '', answers])
            self.gate('review', url)
            digest = sha(folder / '审核包.md')
            self.note(folder, '批准.md', ['审核包哈希: ' + digest, '依据: ' + self.w.say('approve-' + job, '批准 ' + self.urls[job]['title'] + ' 版本 ' + digest[:12])])
        # A substantive change after approval voids it.
        folder = self.w.role_folder(A, self.urls['101']['url'])
        saved = (folder / '审核包.md').read_text(encoding='utf-8')
        (folder / '审核包.md').write_text(saved.replace('第一意向: 合成智能事业群', '第一意向: 合成数据平台部'), encoding='utf-8')
        self.write_target('101', '合成智能事业群', digest=sha(folder / '审核包.md'))
        self.gate('submit', self.urls['101']['url'], allowed=False)
        (folder / '审核包.md').write_text(saved, encoding='utf-8')

    def write_target(self, job, slot, digest=None):
        info = self.urls[job]
        folder = self.w.role_folder(A, info['url'])
        digest = digest or sha(folder / '审核包.md')
        self.note(folder, '目标确认.md', ['公司: ' + A, 'BU: ' + info['bu'], '办公地: ' + info['city'], '岗位: ' + info['title'],
                                       '意向: 第一意向 ' + slot, '数量规则: 每账号最多 2 个，提交后不可改志愿', '审核包哈希: ' + digest,
                                       '依据: ' + self.w.say('target-' + job, '确认投递 ' + info['title'] + ' / ' + slot)])

    def m4_submit(self):
        page, w = self.page, self.w
        for job, slot in (('101', '合成智能事业群'), ('102', '合成数据平台部')):
            info = self.urls[job]
            folder = w.role_folder(A, info['url'])
            self.write_target(job, slot)
            self.gate('submit', info['url'], allowed=False)  # 提交前复核 missing
            page.goto(self.base + '/center')
            self.note(folder, '提交前复核.md', ['复核时间: ' + now(), '账号: synthetic-user', '已投: ' + page.locator('#history').inner_text().replace('\n', ' '),
                                            '审核包哈希: ' + sha(folder / '审核包.md'), '材料: 简历.pdf'])
            # Discard the fill cache after review; independent evidence remains valid.
            excel = sha(w.book)
            (folder / fill_plan.PLAN_NAME).unlink()
            self.gate('review', info['url'])
            self.assertEqual(self.hook('点击“提交申请”'), 2)  # no pass yet: the platform hook blocks
            result = self.cli_gate('submit', info['url'])
            self.assertEqual(result.returncode, 0, result.stdout)
            page.goto(self.base + '/apply/' + job); page.wait_for_selector('#submit')
            self.assertEqual(self.hook('点击“提交申请”'), 0)
            self.gate('submit', info['url'], allowed=False)  # marker blocks a new pass even without JSON
            self.assertEqual(sha(w.book), excel)  # even the click marker never marks Excel submitted
            page.get_by_role('button', name='提交申请').click()
            modal = page.get_by_role('dialog', name='确认投递')
            self.assertIn(info['title'], modal.inner_text()); self.assertIn(slot, modal.inner_text())
            self.assertEqual(self.hook('点击弹窗“确认投递”'), 0)
            page.get_by_role('button', name='确认投递').click()
            page.wait_for_url('**/receipt/*')
            app_id = page.locator('#app-id').inner_text()
            self.shot(folder, '回执.png')
            page.goto(self.base + '/center')
            self.assertIn(app_id, page.locator('#history').inner_text())
            self.note(folder, '申请中心读回.md', ['申请编号: ' + app_id, '申请中心: ' + page.locator('#history').inner_text().replace('\n', ' ')])
            self.assertEqual(self.hook('点击“提交申请”'), 2)  # pass void once the receipt is in the folder
            # Book the receipt in Excel at once and read it back.
            self.gate('prepare', self.urls['102']['url'] if job == '101' else info['url'], allowed=False)
            self.write_tracker(folder / '追踪表读回.md', ['中国大陆|岗位链接|' + info['url'] + '|投递状态=已提交'],
                               lambda wb, u=info['url']: [setattr(r[9], 'value', '已提交') for r in wb['中国大陆'].iter_rows(min_row=2) if r[8].value == u])
            self.assertEqual(Batch(w.book, 'e2e').consistency_errors(), [])

    def m5_resume(self):
        # A new turn: site history, receipts and the tracker must agree.
        self.page.goto(self.base + '/center')
        history = self.page.locator('#history').inner_text()
        rows = {r['岗位链接']: r['投递状态'] for r in Workbook(self.w.book).roles_of(A)}
        self.assertEqual(len(self.site.history), 2)
        self.assertEqual(set(rows.values()), {'已提交'})
        self.note(self.w.bdir / '批次', '续跑读回.md', ['申请中心: ' + history.replace('\n', ' '), '追踪表: 2 个已提交', '冲突: 无'])
        self.gate('submit', self.urls['101']['url'], allowed=False)  # no duplicate submission

    def m6_closeout(self):
        result = self.gate('closeout', allowed=False)  # 收尾.md missing
        self.assertTrue(any('保留页面' in r for r in result['收尾']['原因']))
        self.w.closeout_file(retained='合成公司申请中心（已投 2 个）')
        result = self.gate('closeout')
        report = result['报告']
        self.assertEqual(report['盘点状态'], {'已有相关岗位记录': 1, '无法访问校招页面': 1})
        self.assertEqual(report['留空公司'], {'全部': 1})
        self.assertEqual(report['已提交且有回执'], 2)
        self.log.append({'final_report': report})


if __name__ == '__main__':
    unittest.main(argv=[sys.argv[0]])
