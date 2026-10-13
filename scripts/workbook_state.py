"""Read the role tracker (Excel) as the single source of state.

Schema and labels come from campus-job-skills `references/workbook-schema.md`:
  * regional role sheets (e.g. 中国大陆 / 香港) with 公司 … 岗位链接 投递状态 … 备注;
  * `校招计划` with one row per company–plan and `盘点状态`, one of five fixed values.
This module only reads. Writing stays with the Agent's spreadsheet tool, followed by
`tracker_readback.py`, so user edits, formulas and validation are preserved.
"""
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode

INVENTORY_STATUSES = ('已有相关岗位记录', '无当年校招计划', '无用户要求岗位', '无法访问校招页面', '尚未推进')
SCANNED_NO_ROLE = ('无当年校招计划', '无用户要求岗位', '无法访问校招页面')  # scanned, nothing to apply for: kept gaps
NOT_STARTED = '尚未推进'
APPLY_STATUSES = ('待投递', '已提交')
ROLE_HEADER = ('公司', '岗位名称', '岗位链接', '投递状态')
PLAN_SHEET = '校招计划'


def text(value):
    return value.strip() if isinstance(value, str) else ('' if value is None else str(value).strip())


def canonical_url(value):
    """Direct-link identity: drop tracking parameters only (utm_*, ref, source); keep job/BU/plan parameters."""
    value = text(value)
    if not value:
        return ''
    try:
        parts = urlsplit(value)
    except ValueError:
        return ''
    if parts.scheme not in ('http', 'https') or not parts.netloc:
        return ''
    query = [(k, v) for k, v in parse_qsl(parts.query, keep_blank_values=True)
             if not k.lower().startswith('utm_') and k.lower() not in ('ref', 'source')]
    return urlunsplit((parts.scheme.lower(), parts.netloc.lower(), parts.path, urlencode(sorted(query)), parts.fragment))


class Workbook:
    def __init__(self, path):
        from openpyxl import load_workbook
        self.path = Path(path)
        book = load_workbook(self.path, data_only=True)
        self.roles, self.placeholders, self.plans, self.problems = [], [], [], []
        self.role_sheets = []
        for ws in book.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            header = [text(h) for h in rows[0]]
            if ws.title == PLAN_SHEET:
                self._plans(ws.title, header, rows)
            elif all(h in header for h in ROLE_HEADER):
                self.role_sheets.append(ws.title)
                self._roles(ws.title, header, rows)
        if PLAN_SHEET not in book.sheetnames:
            self.problems.append('工作簿缺少「校招计划」表（盘点状态写在这里）')
        if not self.role_sheets:
            self.problems.append('工作簿没有含 公司/岗位名称/岗位链接/投递状态 的岗位表')

    def _roles(self, sheet, header, rows):
        for number, row in enumerate(rows[1:], start=2):
            record = {h: text(v) for h, v in zip(header, row) if h}
            if not record.get('公司'):
                continue
            record['_sheet'], record['_row'] = sheet, number
            record['_url'] = canonical_url(record.get('岗位链接'))
            (self.roles if record.get('岗位链接') else self.placeholders).append(record)

    def _plans(self, sheet, header, rows):
        if '盘点状态' not in header or '公司' not in header:
            self.problems.append('「校招计划」表缺少 公司 或 盘点状态 列')
            return
        for number, row in enumerate(rows[1:], start=2):
            record = {h: text(v) for h, v in zip(header, row) if h}
            if record.get('公司'):
                record['_sheet'], record['_row'] = sheet, number
                self.plans.append(record)

    # ------------------------------------------------------------ queries
    def companies(self):
        names = [r['公司'] for r in self.roles + self.placeholders] + [p['公司'] for p in self.plans]
        return list(dict.fromkeys(names))

    def plan_statuses(self, company):
        return [p.get('盘点状态', '') for p in self.plans if p['公司'] == company]

    def company_status(self, company):
        """One effective inventory label per company, aggregated over its plan rows (campus rule).

        Any plan with relevant roles wins; otherwise a no-role conclusion; 尚未推进 only when every
        row still says so; '' when the company has no 校招计划 row at all.
        """
        statuses = [s for s in self.plan_statuses(company) if s]
        if not statuses:
            return ''
        for value in ('已有相关岗位记录', *SCANNED_NO_ROLE):
            if value in statuses:
                return value
        return statuses[0]

    def roles_of(self, company):
        return [r for r in self.roles if r['公司'] == company]

    def role_by_url(self, url):
        key = canonical_url(url)
        found = [r for r in self.roles if key and r['_url'] == key]
        return found

    def label_errors(self, companies=None):
        """Structural label checks from workbook-schema.md for the given companies (default: all)."""
        errors = list(self.problems)
        scope = set(companies) if companies is not None else None
        for p in self.plans:
            if scope is not None and p['公司'] not in scope:
                continue
            if p.get('盘点状态') and p['盘点状态'] not in INVENTORY_STATUSES:
                errors.append('「校招计划」第%d行 %s 的盘点状态「%s」不是五个固定值之一' % (p['_row'], p['公司'], p['盘点状态']))
        for r in self.roles:
            if scope is not None and r['公司'] not in scope:
                continue
            where = '%s第%d行 %s' % (r['_sheet'], r['_row'], r['公司'])
            if r.get('投递状态') not in APPLY_STATUSES:
                errors.append(where + ' 投递状态应为 待投递 或 已提交，现为「%s」' % r.get('投递状态', ''))
            if not r['_url']:
                errors.append(where + ' 岗位链接不是有效的官方直链')
        seen = {}
        for r in self.roles:
            if r['_url'] and (scope is None or r['公司'] in scope):
                if r['_url'] in seen:
                    errors.append('岗位链接重复：%s 与 %s第%d行' % (seen[r['_url']], r['_sheet'], r['_row']))
                seen.setdefault(r['_url'], '%s第%d行' % (r['_sheet'], r['_row']))
        for r in self.placeholders:
            if (scope is None or r['公司'] in scope) and r.get('投递状态'):
                errors.append('%s第%d行 %s 是公司占位行，投递状态应为空' % (r['_sheet'], r['_row'], r['公司']))
        for company in (scope if scope is not None else self.companies()):
            status = self.company_status(company)
            if status == '已有相关岗位记录' and not self.roles_of(company):
                errors.append(company + ' 盘点状态为 已有相关岗位记录，但岗位表里没有带直链的岗位行')
        return errors
