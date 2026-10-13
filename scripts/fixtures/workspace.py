"""Synthetic application workspace for tests: the campus sample tracker + 投递证据/ (no personal data).

`Workspace` writes files exactly where the Skill says they go, so tests exercise the same layout
an Agent produces. Company names used for roles are synthetic (合成公司…).
"""
import hashlib
import shutil
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
SAMPLE = HERE / 'sample_tracker.xlsx'
ROLE_SHEET = '中国大陆'


class Workspace:
    def __init__(self, root, batch='b1', companies=(('合成公司A', '大厂'), ('合成公司B', '银行')),
                 targets='大厂=1, 银行=1', rule='补位', basis='有合格岗位的公司'):
        self.root = Path(root)
        self.batch = batch
        self.companies = list(companies)
        self.book = self.root / '岗位追踪表.xlsx'
        shutil.copy(SAMPLE, self.book)
        self.bdir = self.root / '投递证据' / batch
        (self.bdir / '批次').mkdir(parents=True)
        (self.root / '简历.pdf').write_bytes(b'%PDF synthetic resume')
        wb = load_workbook(self.book)
        for company, _ in self.companies:
            wb[ROLE_SHEET].append(['北京', company, '合成行业', '普通私企'])
            wb['校招计划'].append([company, None, None, '中国大陆', None, None, None, None, None, '尚未推进'])
        wb.save(self.book)
        self.write_query(targets=targets, rule=rule, basis=basis)

    # ------------------------------------------------------------ files
    def say(self, name, text='synthetic: 用户确认'):
        """A user confirmation stored as evidence; returns its path relative to the root."""
        path = self.bdir / '确认记录' / (name + '.md')
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text('# 用户确认\n\n原话: ' + text + '\n时间: 2026-10-13T10:00:00+08:00\n', encoding='utf-8')
        return path.relative_to(self.root).as_posix()

    def write(self, rel, text):
        path = self.root / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding='utf-8')
        return path

    def write_query(self, targets, rule, basis, **override):
        sections = {
            '入口': '类型: 投递',
            '硬范围': '地区: 中国大陆\n届别: 2027届\n用工: 正式\n职责: AI 产品',
            '偏好轴': '用工路径: 不限\n城市: 不限\n方向: AI 产品优先\n实习政策: 不收实习\n冲突策略: 用户决定',
            '默认岗位次序': 'AI 产品 > 数据产品',
            '公司清单': '\n'.join('- ' + c + ' | ' + b for c, b in self.companies),
            '目标与停止条件': '计数口径: ' + basis + '\n补位规则: ' + rule + '\n目标: ' + targets + '\n停止条件: 确认清单全部覆盖',
            '材料': '- 简历 | 简历.pdf | ' + hashlib.sha256((self.root / '简历.pdf').read_bytes()).hexdigest(),
            '授权': '选岗授权: 是\n准备授权: 是',
            '执行方式': '执行方式: 一个AI',
            '自动化授权': '模式: 不开启\n已开启: 无\n条款方式: 逐项查看\n验证码: 用户提供',
            '能力预检': '内置浏览器: 可用\nChrome: 可用\ncomputer use: 未检查\n系统脚本: 不适用\nPlaywright: 未检查\nPlaywright代替: 未问',
        }
        sections.update(override)
        body = '批次: ' + self.batch + '\n'
        for name, text in sections.items():
            if text is None:
                continue
            body += '\n## ' + name + '\n' + text + '\n依据: ' + self.say('query-' + name) + '\n'
        self.write('投递证据/' + self.batch + '/批次/Query.md', body)

    # ------------------------------------------------------------ workbook edits (what the Agent writes)
    def edit(self, fn):
        wb = load_workbook(self.book)
        fn(wb)
        wb.save(self.book)

    def tag(self, company, status, note=None):
        def fn(wb):
            for row in wb['校招计划'].iter_rows(min_row=2):
                if row[0].value == company:
                    row[9].value = status
            if note:
                for row in wb[ROLE_SHEET].iter_rows(min_row=2):
                    if row[1].value == company and not row[4].value:
                        row[14].value = note
        self.edit(fn)

    def add_role(self, company, url, title='合成AI产品经理', bu='合成事业群', city='北京', jd='职责：AI 产品。要求：本科及以上。'):
        def fn(wb):
            ws = wb[ROLE_SHEET]
            for row in ws.iter_rows(min_row=2):
                if row[1].value == company and not row[8].value:  # fill the placeholder first
                    for i, v in ((4, title), (5, city), (6, jd), (7, bu), (8, url), (9, '待投递')):
                        row[i].value = v
                    return
            last = max(r[0].row for r in ws.iter_rows(min_row=2) if r[1].value == company)
            ws.insert_rows(last + 1)  # additional roles go directly under the company's rows
            for i, v in enumerate(['北京', company, '合成行业', '普通私企', title, city, jd, bu, url, '待投递']):
                ws.cell(row=last + 1, column=i + 1, value=v)
        self.edit(fn)

    def set_apply_status(self, url, status):
        def fn(wb):
            for row in wb[ROLE_SHEET].iter_rows(min_row=2):
                if row[8].value == url:
                    row[9].value = status
        self.edit(fn)

    # ------------------------------------------------------------ evidence per stage
    def company_ready(self, company, limit=2, state='已核实', lane='内置浏览器'):
        base = '投递证据/%s/公司/%s/' % (self.batch, company)
        rule = '状态: ' + state + ('\n上限: %d\n来源: https://example.test/faq\n核验日期: 2026-10-13' % limit if state == '已核实' else '\n上限: 无')
        self.write(base + '数量规则.md', rule + '\n')
        self.write(base + '次序确认.md', '决定: 不变\n依据: ' + self.say('order-' + company) + '\n')
        self.write(base + '通道.md', '通道: ' + lane + '\n')

    def role_folder(self, company, url):
        import sys
        sys.path.insert(0, str(HERE.parent))
        from evidence_layout import role_dir
        return role_dir(self.root, self.batch, company, url)

    def role_ready(self, company, url, conclusion='符合'):
        folder = self.role_folder(company, url)
        folder.mkdir(parents=True, exist_ok=True)
        (folder / '岗位包.md').write_text('公司: ' + company + '\n直链: ' + url + '\n', encoding='utf-8')
        (folder / '资格核对.md').write_text('结论: ' + conclusion + '\n', encoding='utf-8')
        return folder

    def role_prepared(self, company, url, answer='synthetic answer'):
        folder = self.role_ready(company, url)
        (folder / '填写读回.md').write_text('姓名: 合成候选人\n', encoding='utf-8')
        (folder / '草稿读回.md').write_text('重开后一致: 是\n', encoding='utf-8')
        (folder / '审核包.md').write_text('# 审核包\n关键回答: ' + answer + '\n', encoding='utf-8')
        return folder

    def review_hash(self, company, url):
        return hashlib.sha256((self.role_folder(company, url) / '审核包.md').read_bytes()).hexdigest()

    def role_approved(self, company, url, title='合成AI产品经理', bu='合成事业群', city='北京'):
        folder = self.role_folder(company, url)
        digest = self.review_hash(company, url)
        (folder / '批准.md').write_text('审核包哈希: ' + digest + '\n依据: ' + self.say('approve-' + folder.name) + '\n', encoding='utf-8')
        (folder / '目标确认.md').write_text(
            '公司: %s\nBU: %s\n办公地: %s\n岗位: %s\n意向: 无\n数量规则: 上限内\n审核包哈希: %s\n依据: %s\n'
            % (company, bu, city, title, digest, self.say('target-' + folder.name)), encoding='utf-8')
        (folder / '提交前复核.md').write_text('复核: 岗位、名额、材料、版本一致\n', encoding='utf-8')
        return folder

    def role_submitted(self, company, url, readback=True):
        folder = self.role_folder(company, url)
        (folder / '回执.md').write_text('申请编号: APP-SYNTH\n', encoding='utf-8')
        self.set_apply_status(url, '已提交')
        if readback:
            (folder / '追踪表读回.md').write_text('结果: 通过\n', encoding='utf-8')
        return folder

    def closeout_file(self, retained='无（全部回执页已截图存档）'):
        self.write('投递证据/%s/批次/收尾.md' % self.batch,
                   '停止原因: 确认清单已全部覆盖\n保留页面: ' + retained + '\n依据: ' + self.say('stop') + '\n')
