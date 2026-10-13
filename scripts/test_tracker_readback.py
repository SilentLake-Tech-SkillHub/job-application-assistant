#!/usr/bin/env python3
"""4.6 / #25: reopen the saved tracker and compare with the pre-write backup (synthetic rows on the sample workbook)."""
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from tracker_readback import check, parse_expect

HERE = Path(__file__).resolve().parent
URL = 'https://example.test/jobs/101'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class TrackerReadbackTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); d = Path(self.tmp.name)
        self.book, self.backup = d / 'tracker.xlsx', d / 'tracker.bak.xlsx'
        shutil.copy(HERE / 'fixtures' / 'sample_tracker.xlsx', self.book)
        wb = load_workbook(self.book)
        wb['中国大陆'].append(['北京', '合成公司A', '合成', '普通私企', '合成AI产品经理', '北京', 'jd', 'bu', URL, '待投递'])
        wb['校招计划'].append(['合成公司A', None, None, '中国大陆', None, None, None, None, None, '尚未推进'])
        wb.save(self.book)
        shutil.copy(self.book, self.backup); self.pre = sha(self.backup)
        self.expect = [parse_expect('中国大陆|岗位链接|' + URL + '|投递状态=已提交'),
                       parse_expect('校招计划|公司|合成公司A|盘点状态=已有相关岗位记录')]

    def tearDown(self):
        self.tmp.cleanup()

    def write(self, mutate):
        wb = load_workbook(self.book); mutate(wb); wb.save(self.book)

    def row(self, wb):
        return next(r for r in wb['中国大陆'].iter_rows(min_row=2) if r[8].value == URL)

    def plan(self, wb):
        return next(r for r in wb['校招计划'].iter_rows(min_row=2) if r[0].value == '合成公司A')

    def good(self, wb):
        self.row(wb)[9].value = '已提交'
        self.plan(wb)[9].value = '已有相关岗位记录'

    def test_both_labels_written_and_read_back(self):
        self.write(self.good)
        self.assertEqual(check(self.book, self.backup, self.pre, self.expect), [])

    def test_unsaved_write_is_detected(self):
        errors = check(self.book, self.backup, self.pre, self.expect)
        self.assertTrue(any('byte-identical' in e for e in errors))

    def test_missing_label_write_is_detected(self):
        self.write(lambda wb: setattr(self.row(wb)[9], 'value', '已提交'))
        self.assertTrue(any("盘点状态 reads back '尚未推进'" in e for e in check(self.book, self.backup, self.pre, self.expect)))

    def test_unrelated_rows_and_sheets_are_protected(self):
        def mutate(wb):
            self.good(wb)
            wb['中国大陆']['B2'] = '被改了的公司'
            wb['分类说明']['B2'] = '被覆盖'
        self.write(mutate)
        errors = check(self.book, self.backup, self.pre, self.expect)
        self.assertTrue(any('other sheet changed: 分类说明' in e for e in errors))
        self.assertTrue(any('placeholder row changed or removed' in e for e in errors))

    def test_placeholder_may_become_first_role_row(self):
        wb = load_workbook(self.book)
        wb['中国大陆'].append(['北京', '合成公司B', '合成', '普通私企'])
        wb.save(self.book); shutil.copy(self.book, self.backup); self.pre = sha(self.backup)
        url = 'https://example.test/jobs/201'
        def mutate(wb):
            row = next(r for r in wb['中国大陆'].iter_rows(min_row=2) if r[1].value == '合成公司B')
            for i, v in ((4, '合成岗位'), (8, url), (9, '待投递')):
                row[i].value = v
        self.write(mutate)
        self.assertEqual(check(self.book, self.backup, self.pre, [parse_expect('中国大陆|岗位链接|' + url + '|投递状态=待投递')]), [])

    def test_duplicate_row_error_values_hash_and_backup(self):
        def mutate(wb):
            self.good(wb)
            wb['中国大陆'].append(['北京', '合成公司A', '合成', '普通私企', '重复', '北京', 'jd', 'bu', URL, '已提交'])
            wb['香港']['P2'] = '#REF!'
        self.write(mutate)
        errors = check(self.book, self.backup, self.pre, self.expect)
        self.assertTrue(any('expected exactly one row' in e for e in errors))
        self.assertTrue(any('#REF!' in e for e in errors))
        self.assertTrue(any('backup SHA-256' in e for e in check(self.book, self.backup, '0' * 64, self.expect)))
        self.backup.unlink()
        self.assertEqual(check(self.book, self.backup, self.pre, self.expect), ['backup not found: make a recoverable backup before writing'])

    def test_cli_records_markdown_evidence(self):
        self.write(self.good)
        record = Path(self.tmp.name) / '投递证据' / 'b1' / '岗位' / 'x' / '追踪表读回.md'
        args = [sys.executable, str(HERE / 'tracker_readback.py'), '--workbook', str(self.book), '--backup', str(self.backup),
                '--pre-sha256', self.pre, '--expect', '中国大陆|岗位链接|' + URL + '|投递状态=已提交',
                '--expect', '校招计划|公司|合成公司A|盘点状态=已有相关岗位记录', '--record', str(record)]
        result = subprocess.run(args, capture_output=True, text=True, encoding='utf-8', env=dict(os.environ, PYTHONIOENCODING='utf-8'))
        self.assertEqual(result.returncode, 0, result.stdout)
        self.assertTrue(json.loads(result.stdout)['读回通过'])
        self.assertIn('结果: 通过', record.read_text(encoding='utf-8'))

    def test_bad_expect_format(self):
        with self.assertRaises(ValueError):
            parse_expect('only|three|parts')


if __name__ == '__main__':
    unittest.main()
