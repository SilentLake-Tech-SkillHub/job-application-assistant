#!/usr/bin/env python3
"""Shared fixed writer and complete business-data snapshot for recruiting workbooks."""
import argparse
import copy
import json
import os
import sys
from datetime import date, datetime, time
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, get_column_letter

FULL_NAME = '工作簿全量.json'
FULL_SCHEMA = 2


def value(v):
    if isinstance(v, (datetime, date, time)):
        return {'type': type(v).__name__, 'value': v.isoformat()}
    return v


def snapshot(book):
    """All sheets/cells, including formulas, cached values, notes and real links; no styles."""
    from company_handoff import sha, digest, read, SCHEMA
    before = sha(book)
    wb = load_workbook(book, data_only=False)
    cached = load_workbook(book, data_only=True)
    sheets = []
    try:
        for ws in wb:
            rows = []
            for row in ws:
                cells = []
                for cell in row:
                    if cell.value is None and not getattr(cell,'hyperlink',None) and not getattr(cell,'comment',None):
                        continue
                    item = {'column': cell.column, 'value': value(cell.value), 'type': cell.data_type}
                    if cell.data_type == 'f':
                        item['cached_value'] = value(cached[ws.title][cell.coordinate].value)
                    if cell.hyperlink:
                        item['hyperlink'] = {'target': cell.hyperlink.target, 'location': cell.hyperlink.location}
                    if cell.comment:
                        item['note'] = {'text':cell.comment.text,'author':cell.comment.author}
                    cells.append(item)
                if cells:
                    rows.append({'row': row[0].row, 'cells': cells})
            sheets.append({'name': ws.title, 'rows': rows, 'max_row': ws.max_row,
                           'max_column': ws.max_column, 'state': ws.sheet_state,
                           'merged_ranges': [str(x) for x in ws.merged_cells.ranges],
                           'tables': [{'name': t.name, 'ref': t.ref} for t in ws.tables.values()]})
    finally:
        wb.close(); cached.close()
    batches={}
    shared={s['name'] for s in sheets}
    if {s[0] for s in SCHEMA.values()}<=shared:
        company_sheet=next(s for s in sheets if s['name']=='公司核对')
        ids={next((c['value'] for c in row['cells'] if c['column']==1),None) for row in company_sheet['rows'] if row['row']>1}
        for batch in sorted(x for x in ids if x):batches[batch]=read(book,batch)
    if before != sha(book):
        raise ValueError('工作簿在完整导出期间变化')
    return {'schema_version': FULL_SCHEMA, 'scope': 'all_workbook_business_data',
            'workbook_name': Path(book).name, 'workbook_sha256': before,
            'sheets': sheets, 'company_batches':batches, 'data_sha256': digest({'sheets':sheets,'company_batches':batches})}


def full_path(book):
    return Path(book).resolve().parent / '岗位交接' / FULL_NAME


def export_full(book):
    from company_handoff import atomic, sha
    data = snapshot(book); out = full_path(book)
    if out.exists():
        atomic(out.parent / '历史' / (sha(out) + '.json'), out.read_bytes())
    atomic(out, json.dumps(data, ensure_ascii=False, indent=2).encode('utf-8'))
    # Compare the serialized data, not just the filename or write return code.
    if json.loads(out.read_text(encoding='utf-8')) != data:
        raise ValueError('完整JSON读回不一致')
    return data


def check_full(book):
    try:
        saved = json.loads(full_path(book).read_text(encoding='utf-8'))
        return [] if saved == snapshot(book) else ['完整JSON与当前Excel业务数据不一致']
    except (OSError, ValueError, TypeError) as exc:
        return ['完整JSON不可读取: ' + str(exc)]


def preserve_empty_text(path, coordinates):
    """openpyxl drops empty string payloads on save; keep this distinct source value."""
    from zipfile import ZipFile, ZIP_DEFLATED
    from xml.etree import ElementTree as ET
    ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    relns='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    with ZipFile(path) as z:
        files={item.filename:(item,z.read(item.filename)) for item in z.infolist()}
    workbook=ET.fromstring(files['xl/workbook.xml'][1])
    rels=ET.fromstring(files['xl/_rels/workbook.xml.rels'][1])
    targets={r.attrib['Id']:r.attrib['Target'] for r in rels}
    for sheet in workbook.find('{'+ns+'}sheets'):
        wanted=coordinates.get(sheet.attrib['name'],set())
        if not wanted:continue
        target=targets[sheet.attrib['{'+relns+'}id']]
        name=target.lstrip('/') if target.startswith('/') else 'xl/'+target
        info,raw=files[name]
        # Preserve namespace spellings and extensions in the rest of the sheet.
        # Only replace the targeted empty inline cell element, not serialize the full tree.
        import re
        def restore(match):
            cell=match.group(0);coordinate=re.search(rb'\br="([^"]+)"',cell)
            if not coordinate or coordinate.group(1).decode() not in wanted or b'<is' in cell:return cell
            if cell.endswith(b'/>'):return cell[:-2]+b'><is><t></t></is></c>'
            return cell[:-4]+b'<is><t></t></is></c>'
        raw=re.sub(rb'<c\b[^>]*?(?:/>|>.*?</c>)',restore,raw,flags=re.DOTALL)
        files[name]=(info,raw)
    temporary=Path(path).with_suffix('.zip.tmp')
    try:
        with ZipFile(temporary,'w',ZIP_DEFLATED) as z:
            for info,raw in files.values():z.writestr(info,raw)
        os.replace(temporary,path)
    finally:temporary.unlink(missing_ok=True)


def apply_updates(wb, updates, root):
    """Explicit keyed upserts: no entire-table replacement, silent fields or status reset."""
    from company_handoff import safe_ref, SCHEMA
    for update in updates:
        if set(update) - {'sheet', 'key', 'values', 'receipt_ref'}:
            raise ValueError('工作表更新含未知字段')
        sheet, key, fields = update['sheet'], update['key'], update['values']
        if sheet in {v[0] for v in SCHEMA.values()}:
            raise ValueError('共享表须通过companies更新')
        if sheet not in wb or not key or not fields:
            raise ValueError('更新需要既有工作表、非空复合键和字段')
        ws = wb[sheet]; heads = [c.value for c in ws[1]]
        if len(heads) != len(set(heads)) or None in heads:
            raise ValueError('固定写入要求唯一非空表头: ' + sheet)
        if not (set(key) | set(fields)) <= set(heads):
            raise ValueError('未知表头: ' + sheet)
        if any(v in ('', None) for v in key.values()):
            raise ValueError('禁止空身份键')
        if any(k in fields and fields[k] != v for k, v in key.items()):
            raise ValueError('不可通过字段修改记录身份')
        indexes = {h: i + 1 for i, h in enumerate(heads)}
        matches = [row[0].row for row in ws.iter_rows(min_row=2)
                   if all(row[indexes[k] - 1].value == v for k, v in key.items())]
        if len(matches) > 1:
            raise ValueError('复合键匹配多行，须补齐身份: ' + sheet)
        old = {h: ws.cell(matches[0], i).value for h, i in indexes.items()} if matches else {}
        if old.get('投递状态') == '已提交' and fields.get('投递状态', '已提交') != '已提交':
            raise ValueError('禁止重置已有提交')
        if fields.get('投递状态') == '已提交' and old.get('投递状态') != '已提交':
            safe_ref(root, update.get('receipt_ref', ''))
        if '投递状态' in fields and fields['投递状态'] not in ('', '待投递', '已提交'):
            raise ValueError('非法投递状态')
        if matches:
            number = matches[0]
        else:
            company = fields.get('公司', key.get('公司'))
            same = [r[0].row for r in ws.iter_rows(min_row=2)
                    if company and '公司' in indexes and r[indexes['公司'] - 1].value == company]
            placeholders = [n for n in same if '岗位链接' in indexes and not ws.cell(n, indexes['岗位链接']).value
                            and not ws.cell(n, indexes.get('岗位名称', indexes['岗位链接'])).value]
            if placeholders and (fields.get('岗位链接') or key.get('岗位链接')):
                number = placeholders[0]
            else:
                number = same[-1] + 1 if same else ws.max_row + 1
                # Insertion would shift arbitrary formulas without translating dependencies.
                if number <= ws.max_row and any(c.data_type == 'f' for row in ws for c in row):
                    raise ValueError('含公式表插行须专门处理引用，禁止静默移动')
                ws.insert_rows(number)
                if same:
                    template = same[0] if same[0] < number else same[0] + 1
                    for column in range(1, ws.max_column + 1):
                        ws.cell(number, column)._style = copy.copy(ws.cell(template, column)._style)
                    for field in ('地区','公司','主要行业','企业性质'):
                        if field in indexes:
                            ws.cell(number, indexes[field]).value = ws.cell(template, indexes[field]).value
                for table in ws.tables.values():
                    a, b, c, d = range_boundaries(table.ref)
                    if b < number <= d + 1:
                        table.ref = f'{get_column_letter(a)}{b}:{get_column_letter(c)}{d + 1}'
                        if table.autoFilter:
                            table.autoFilter.ref = table.ref
        for k, v in {**key, **fields}.items():
            if isinstance(v,dict) and set(v)=={'type','value'} and v['type'] in ('date','datetime','time'):
                v={'date':date,'datetime':datetime,'time':time}[v['type']].fromisoformat(v['value'])
            if isinstance(v, (list, dict)):
                raise ValueError('原表字段需标量；结构化明细写对应共享表')
            if isinstance(v, str) and len(v) > 32767:
                raise ValueError('禁止截断单元格')
            cell = ws.cell(number, indexes[k])
            if cell.data_type == 'f' and cell.value != v:
                raise ValueError('禁止覆盖已有公式')
            cell.value = v
            if isinstance(v,(date,datetime)) and cell.number_format=='General':cell.number_format='yyyy-mm-dd'
            if isinstance(v, str):
                cell.data_type = 's'
        if ws.auto_filter.ref:
            ws.auto_filter.ref = ws.dimensions


def refresh(book, batch=None):
    from company_handoff import export_batch
    full = export_full(book)
    batches = set()
    wb = load_workbook(book, read_only=True)
    try:
        if '公司核对' in wb:
            batches.update(r[0] for r in wb['公司核对'].iter_rows(min_row=2, values_only=True) if r[0])
    finally:
        wb.close()
    if batch:
        batches.add(batch)
    for item in sorted(batches):
        export_batch(book, item, full)
    from company_handoff import sha
    if sha(book)!=full['workbook_sha256']:raise ValueError('刷新期间工作簿变化')
    return {'workbook_sha256': full['workbook_sha256'], 'batches': sorted(batches), 'sheets': len(full['sheets'])}


def query(book, company=None, sheet=None):
    """Token-saving reads from the verified complete snapshot, with original columns retained."""
    from company_handoff import sha, digest
    data = json.loads(full_path(book).read_text(encoding='utf-8'))
    if data['workbook_sha256'] != sha(book) or data['data_sha256'] != digest({'sheets':data['sheets'],'company_batches':data.get('company_batches',{})}):
        raise ValueError('完整JSON已陈旧，先refresh')
    result = []
    company_keys=set()
    for ws in data['sheets']:
        if ws['name']=='公司核对':
            for row in ws['rows'][1:]:
                cells={c['column']:c['value'] for c in row['cells']}
                if cells.get(3)==company:company_keys.add(cells.get(2))
    for ws in data['sheets']:
        if sheet and ws['name'] != sheet:
            continue
        rows = ws['rows']
        header = next((r for r in rows if r['row'] == 1), {'cells': []})
        names = {c['column']: c['value'] for c in header['cells'] if isinstance(c['value'], str)}
        for row in rows:
            if row['row'] == 1:
                continue
            fields = {names.get(c['column'], str(c['column'])): c['value'] for c in row['cells']}
            if company and fields.get('公司') != company and fields.get('公司键') not in company_keys:
                continue
            result.append({'sheet': ws['name'], 'row': row['row'], 'values': fields})
    return result



def pending_tasks():
    from company_handoff import RESEARCH_ITEMS
    return [{'id':item,'status':'todo','source':'','checked_at':'','gap':'须按当前范围查证；历史记录只作线索','expected_count':'','observed_count':''} for item in sorted(RESEARCH_ITEMS)]


def initialize(book,batch,sheet,scope,company=None):
    """Create explicit research tasks; preserve old facts without asserting new verification."""
    from company_handoff import read,write,DIMENSIONS,digest
    from structured_dimensions import QUOTA_DIMENSIONS
    if not batch or not sheet or not scope:
        raise ValueError('init需要batch、sheet与已确认scope')
    wb=load_workbook(book,read_only=True)
    try:
        ws=wb[sheet];heads=next(ws.values)
        if '公司' not in heads:raise ValueError('指定表没有公司列')
        idx=heads.index('公司')
        names=list(dict.fromkeys(row[idx] for row in ws.iter_rows(min_row=2,values_only=True) if row[idx] and (not company or row[idx]==company)))
        existing=read(book,batch,allow_legacy=True) if '公司核对' in wb else []
    finally:wb.close()
    if company and not names:raise ValueError('指定公司不在工作表')
    known={c['name'] for c in existing};items=[]
    for name in names:
        if name in known:continue
        items.append({'company_key':digest([sheet,name])[:16],'name':name,'version':'draft-v1','scope':scope,
          'scope_status':'unknown','coverage':'partial','gaps':'尚未按新流程核验；历史事实保留在原表及existing_data','checked_at':'',
          'recommendation_reason':'研究未完成，尚未推荐或选择','plans':[],'roles':[],
          'rules':[{'id':dim,'dimension':dim,'scope':'公司当前范围，适用关系待查','status':'unknown','limit_type':'unknown' if dim in QUOTA_DIMENSIONS else 'not_applicable','usage_status':'unknown','text':'待核实','gap':'尚未按当前范围查证'} for dim in sorted(DIMENSIONS)],
          'referrals':[{'id':'pending','scope':'计划/BU适用范围待查','status':'unverified','notes':'须搜索适用内推码；保留原表历史值，不视为当前已核实'}],
          'conditions':[],'preferences':[],'research':pending_tasks()})
    # Upgrade existing v1 companies without turning old selections into new consent.
    for c in existing:
        if not c.get('scope_status') or not c['research'] or 'priority_scope' not in c or ('preference' not in {r['dimension'] for r in c['rules']} and not c['preferences']):
            c['scope_status']='unknown';c['version']+='-schema-v2';c['coverage']='partial';c['gaps']=(c['gaps']+'；新流程检查待补').strip('；')
            c['research']=pending_tasks()
            c.setdefault('priority_scope','unknown')
            c.setdefault('user_priorities',[{'id':'up-'+r['id'],'role_id':r['id'],'scope':'unknown','status':'unknown','rank':None} for r in c['roles']])
            c['preferences']=[{'id':'pref-'+r['id'],'role_id':r['id'],'status':'unknown','mode':'unknown','rank':None} for r in c['roles']]
            if 'preference' not in {r['dimension'] for r in c['rules']}:
                c['rules'].append({'id':'preference','dimension':'preference','scope':'公司当前范围','status':'unknown','text':'志愿规则待核实','gap':'历史协议缺少独立志愿'})
            for rule in c['rules']:
                if not rule.get('limit_type'):
                    rule['limit_type']='count' if type(rule.get('limit'))==int else 'not_applicable' if rule['status']=='not_applicable' else 'unknown'
                    if rule['limit_type']=='unknown':rule.update(status='unknown',gap='原规则文字待拆分；未将空白上限解释为无限制')
                rule.setdefault('usage_status','verified' if type(rule.get('used'))==int else 'unknown')
            items.append(c)
    if items:write(book,batch,{'companies':items})
    else:refresh(book,batch)
    return {'initialized_or_migrated':len(items),'source_companies':len(names),'batch':batch}

def main():
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8', errors='backslashreplace')
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('command', choices=['refresh', 'check', 'query','init'])
    p.add_argument('--workbook', required=True, type=Path)
    p.add_argument('--scope'); p.add_argument('--batch'); p.add_argument('--company'); p.add_argument('--sheet')
    a = p.parse_args()
    try:
        if a.command == 'init':
            result = initialize(a.workbook,a.batch,a.sheet,a.scope,a.company)
        elif a.command == 'refresh':
            result = refresh(a.workbook, a.batch)
        elif a.command == 'check':
            errors = check_full(a.workbook); result = {'passed': not errors, 'errors': errors}
        else:
            result = query(a.workbook, a.company, a.sheet)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return int(a.command == 'check' and bool(errors))
    except (OSError, ValueError, KeyError, TypeError) as exc:
        print(str(exc)); return 1


if __name__ == '__main__':
    raise SystemExit(main())
